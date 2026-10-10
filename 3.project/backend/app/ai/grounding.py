"""Evidence anchors and conservative field/party checks. Never consult gold labels."""
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

GUARD_VERSION = "grounding-v3-claim-context-20261010"

# A space is a thousands separator only in complete groups of three digits.
MONEY_TOKEN = r"(?:\d{1,3}(?:[ ,]\d{3})+|\d+)(?:\.\d+)?"

LABELS = {
    "company_name": r"상\s*호|업체명|예금주|수급인(?:명|\s*\(을\))?|공급자명|계약상대자",
    "business_registration_no": r"사업자\s*등록\s*번호|사업자\s*번호|등록\s*번호",
    "supply_amount": r"공급\s*가액|공급\s*금액",
    "tax_amount": r"부가\s*가치세|부가세|세액",
    "total_amount": r"합계\s*금액|총\s*금액|총액|합계|계약\s*대금|계약\s*금액|견적\s*금액|검수\s*금액|청구\s*금액",
    "issue_date": r"발행\s*일|발급\s*일|작성\s*일",
    "amount_scope": r"전액|전체\s*계약|계약\s*전체|선금|선급금|착수금|중도금|잔금|계약금액의|부분\s*청구|분할|계약서|견적서|검수확인서"
}
BUYER = re.compile(r"공급\s*받는\s*자|매입자|발주자")
SUPPLIER = re.compile(r"공급자|매출자|수급인|계약상대자")


def canonical(field, value):
    if value is None or value == "" or value == "UNKNOWN":
        return None
    token = str(value).strip()
    if field.endswith("amount"):
        try:
            token = token.removesuffix("원").strip()
            if not re.fullmatch(MONEY_TOKEN, token):
                return None
            number = Decimal(token.replace(",", "").replace(" ", ""))
            return str(number.normalize()) if number.is_finite() and number >= 0 else None
        except InvalidOperation:
            return None
    if field == "business_registration_no":
        return re.sub(r"[\s-]", "", token)
    if field == "issue_date":
        m = re.fullmatch(r"(\d{4})\s*[년./-]\s*(\d{1,2})\s*[월./-]\s*(\d{1,2})\s*일?", token)
        if not m:
            return None
        try:
            return date(*map(int, m.groups())).isoformat()
        except ValueError:
            return None
    return re.sub(r"\s+", " ", token)


def docx_locations(path, body):
    """Map the existing flat extraction back to native paragraphs/table cells.

    This verifies text, not Word's physical page layout. A changed/mismatching
    flat extraction has no verified structural locations and fails closed.
    """
    from docx import Document
    document = Document(path)
    parts, spans = [], []
    offset = 0

    def append(text, segments):
        nonlocal offset
        if parts:
            offset += 1  # the extractor's newline
        parts.append(text)
        spans.extend({**position, "start": offset + start, "end": offset + end}
                     for start, end, position in segments)
        offset += len(text)

    for pi, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            append(text, [(0, len(text), {"location_kind": "paragraph", "paragraph": pi})])
    for ti, table in enumerate(document.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            cells, segments, cursor = [], [], 0
            for ci, cell in enumerate(row.cells, 1):
                text = cell.text.strip()
                if not text or cells and cells[-1] == text:
                    continue
                if cells:
                    cursor += 1  # tab between cells, as in preprocessor
                trim = len(cell.text) - len(cell.text.lstrip())
                paragraph_offset = 0
                for pi, paragraph in enumerate(cell.paragraphs, 1):
                    start = max(0, paragraph_offset - trim)
                    end = min(len(text), paragraph_offset + len(paragraph.text) - trim)
                    if end > start:
                        segments.append((cursor + start, cursor + end, {
                            "location_kind": "table_cell", "table": ti, "row": ri,
                            "cell": ci, "paragraph": pi}))
                    paragraph_offset += len(paragraph.text) + 1
                cells.append(text)
                cursor += len(text)
            if cells:
                append("\t".join(cells), segments)
    if "\n".join(parts).strip() != body:
        return {}
    locations, offset = {}, 0
    for li, line in enumerate(body.splitlines(), 1):
        locations[str(li)] = [dict(span) for span in spans
                              if span["start"] < offset + len(line) and span["end"] > offset]
        offset += len(line) + 1
    return locations


def source_pages(file_path, raw_text):
    path = Path(file_path)
    body = raw_text.partition("\n")[2] if raw_text.startswith("[파일명:") else raw_text
    if path.suffix.lower() == ".docx":
        return [{"page": None, "location_kind": "document_block", "text": body,
                 "document_format": "docx", "page_layout_verified": False,
                 "line_locations": docx_locations(path, body)}]
    if path.suffix.lower() != ".pdf":
        physical = path.suffix.lower() in {".png", ".jpg", ".jpeg", ".bmp", ".webp"}
        return [{"page": 1 if physical else None, "location_kind": "page" if physical else "document_block", "text": body}]
    # One pagewise extraction path owns PDF native/OCR coverage and raster
    # budgets. Missing evidence must retain its page slot and completeness flag.
    from .preprocessor import extract_pdf_pages
    pages = extract_pdf_pages(path)
    extracted = "\n\n".join(page["text"] for page in pages if page["text"]).strip()
    if extracted != body.strip():
        # A fresh OCR pass can differ from the rule extraction. Do not apply
        # proposals from that different text to the earlier field snapshot.
        for page in pages:
            page["evidence_complete"] = False
            page["extraction_warnings"] = list(dict.fromkeys(
                [*page.get("extraction_warnings", []), "SOURCE_TEXT_CHANGED"]))
    return pages


def source_evidence_complete(pages):
    """Coverage, not model confidence: unreadable text cannot prove absence.

    Older text-only/evaluation callers have no coverage metadata, but must still
    provide nonempty text. DOCX also needs verified native structural locations.
    """
    readable = any(isinstance(page, dict) and isinstance(page.get("text"), str)
                   and bool(page["text"].strip()) for page in pages)
    return readable and all(
        isinstance(page, dict) and isinstance(page.get("text"), str)
        and (bool(page["text"].strip()) or page.get("verified_blank") is True)
        and page.get("evidence_complete", True) is True
        and (page.get("document_format") != "docx" or bool(page.get("line_locations")))
        for page in pages)


def proposal_preflight(field, proposal, pages):
    from .local_llm import valid_answer
    reasons = []
    try:
        valid = field in LABELS and valid_answer(proposal)
    except (TypeError, ValueError):
        valid = False
    if not valid:
        reasons.append("INVALID_FORMAT")
    if not source_evidence_complete(pages):
        reasons.append("SOURCE_EVIDENCE_UNAVAILABLE")
        if any(isinstance(page, dict) and page.get("document_format") == "docx"
               and not page.get("line_locations") for page in pages):
            reasons.append("DOCX_LOCATION_UNVERIFIED")
    if reasons:
        return {"field": field, "proposal": proposal, "accepted": False,
                "reasons": reasons, "warnings": [], "anchor": None,
                "human_confirmed": False}
    return None


def party_at(lines, index, document_type):
    party = "unknown"
    for line in lines[:index + 1]:
        buyer, supplier = BUYER.search(line), SUPPLIER.search(line)
        if buyer and supplier:
            party = "ambiguous"
        elif buyer:
            party = "buyer"
        elif supplier:
            party = "supplier"
    if party == "unknown" and document_type in {"business_registration", "bank_account", "inspection_confirmation"}:
        party = "supplier"
    return party


def candidates(pages, field, document_type):
    found = []
    for page in pages:
        lines = page["text"].splitlines()
        for index, line in enumerate(lines):
            for label in re.finditer(LABELS[field], line):
                # OCR can render the colon as a dot followed by whitespace.
                # A leading decimal dot without whitespace is never a separator.
                tail = re.sub(r"^\s*(?:[:：=]|[.。•·](?=\s))?\s*", "", line[label.end():])
                value_index = index
                if field != "amount_scope" and not tail and index + 1 < len(lines):
                    tail = lines[index + 1]
                    value_index += 1
                value = None
                if field.endswith("amount"):
                    match = re.match(r"(?:[₩￦]|금)?\s*(" + MONEY_TOKEN + r")(?![\d,])", tail)
                    if match and not re.match(r"\s*(?:[%년월일]|[./-]\s*\d|\d)", tail[match.end():]):
                        value = match.group(1)
                        if " " in value and not (re.match(r"\s*원(?:$|\s|[.,(])", tail[match.end():])
                                                     or re.match(r"[₩￦]", tail)):
                            value = None
                    if field == "total_amount" and re.search(r"부가세\s*(?:별도|미포함)|VAT\s*(?:별도|제외)|공급가액\s*기준", line, re.I):
                        value = None
                elif field == "business_registration_no":
                    match = re.match(r"(\d{3}[\s-]*\d{2}[\s-]*\d{5})(?!\d)", tail)
                    value = match.group(1) if match else None
                elif field == "issue_date":
                    match = re.match(r"(\d{4}\s*[년./-]\s*\d{1,2}\s*[월./-]\s*\d{1,2}\s*일?)", tail)
                    value = match.group(1) if match else None
                elif field == "company_name":
                    value = tail.strip() or None
                else:
                    if re.fullmatch(r"계약서|견적서|검수확인서", label.group()) and not candidates(pages, "total_amount", document_type):
                        continue
                    from ..validation.billing import scope_in_document
                    value = scope_in_document(line, document_type, bool(candidates(pages, 'total_amount', document_type)))
                    if value == 'UNKNOWN':
                        continue
                token = canonical(field, value)
                if token is not None:
                    found.append({"value": token, "page": page["page"], "line": value_index + 1,
                                  "label_line": index + 1, "label": label.group(),
                                  "party": party_at(lines, value_index, document_type)})
    return found


def checksum_valid(number):
    if not re.fullmatch(r"\d{10}", number or ""):
        return False
    digits = list(map(int, number))
    total = sum(n * weight for n, weight in zip(digits[:9], (1, 3, 7, 1, 3, 7, 1, 3, 5))) + digits[8] * 5 // 10
    return (10 - total % 10) % 10 == digits[9]


def evidence_anchor(page, line_no, quote):
    lines = page["text"].splitlines()
    if not isinstance(line_no, int) or not 1 <= line_no <= len(lines) or not quote or quote not in lines[line_no - 1]:
        return None
    start = sum(len(line) + 1 for line in lines[:line_no - 1]) + lines[line_no - 1].index(quote)
    anchor = {"page": page["page"], "line": line_no, "start": start, "end": start + len(quote),
              "quote": quote, "location_kind": page["location_kind"]}
    if page.get("document_format") == "docx":
        locations = [p for p in page.get("line_locations", {}).get(str(line_no), [])
                     if p["start"] < anchor["end"] and p["end"] > start]
        if not locations:
            return None
        anchor.update(location_kind="docx_structure", positions=locations,
                      page_layout_verified=False, text_verified=True)
    elif page["location_kind"] != "page":
        return None
    return anchor


def guard_proposal(field, proposal, pages, document_type):
    rejected = proposal_preflight(field, proposal, pages)
    if rejected:
        return rejected
    reasons, warnings = [], []
    state, value = proposal.get("state"), proposal.get("value")
    token = canonical(field, value)
    options = candidates(pages, field, document_type)
    if field in {"company_name", "business_registration_no"}:
        eligible = [item for item in options if item["party"] == "supplier"]
    else:
        eligible = options
    distinct = {item["value"] for item in eligible}
    if state in {"ABSENT", "REVIEW"}:
        if value not in (None, ""):
            reasons.append("INVALID_FORMAT")
        if state == "ABSENT" and options or state == "REVIEW" and len(distinct) == 1:
            reasons.append("UNSUPPORTED_ABSTENTION")
        if (state == "ABSENT" and field in {"supply_amount", "tax_amount", "total_amount", "amount_scope"}
                and any(re.search(r"참고값\s*[A-ZＡ-Ｚ]?\s*[:：.]?\s*\d", p["text"]) for p in pages)):
            reasons.append("AMBIGUOUS_UNLABELED_AMOUNT")
        return {"field": field, "proposal": proposal, "accepted": not reasons, "reasons": reasons,
                "warnings": warnings, "anchor": None, "human_confirmed": False}
    if state != "VALUE" or token is None:
        reasons.append("INVALID_FORMAT")
    quote, line_no = proposal.get("quote", ""), proposal.get("line")
    page_no = proposal.get("page")
    # Scope is a semantic enum, not literal FULL/PARTIAL text. Recover a missing
    # anchor only from one scope meaning AND one explicitly labeled total. A
    # supplied but wrong quote/position is never replaced by this recovery.
    if (field == "amount_scope" and state == "VALUE" and page_no is None and line_no is None
            and not quote and len(distinct) == 1 and token in distinct
            and len({o["value"] for o in candidates(pages, "total_amount", document_type)}) == 1):
        scope_option = eligible[0]
        page_no, line_no = scope_option["page"], scope_option["line"]
        scope_page = next(p for p in pages if p["page"] == page_no)
        quote = scope_page["text"].splitlines()[line_no - 1]
        warnings.append("CODE_ANCHORED_SCOPE")
    matching_page = next((p for p in pages if p["page"] == page_no), None)
    structural = bool(matching_page and matching_page.get("document_format") == "docx")
    if not matching_page or not structural and (page_no is None or matching_page["location_kind"] != "page"):
        reasons.append("PAGE_UNVERIFIED")
    anchor = None
    if matching_page and isinstance(line_no, int) and line_no >= 1:
        lines = matching_page["text"].splitlines()
        # Exact quote + same field/value candidate within one line, same party.
        # The resolved position is saved separately from the original proposal.
        choices = []
        for resolved in range(max(1, line_no - 1), min(len(lines), line_no + 1) + 1):
            nearby = [o for o in options if o["page"] == page_no and o["line"] == resolved and o["value"] == token]
            if resolved != line_no and (not nearby or party_at(lines, min(line_no, len(lines)) - 1, document_type)
                                        != party_at(lines, resolved - 1, document_type)):
                continue
            candidate_anchor = evidence_anchor(matching_page, resolved, quote)
            if candidate_anchor is not None:
                choices.append(candidate_anchor)
        exact = [a for a in choices if a["line"] == line_no]
        if exact:
            anchor = exact[0]
        elif len(choices) == 1:
            anchor = choices[0]
            warnings.append("LINE_OFFSET_RESOLVED")
        if anchor:
            line_no = anchor["line"]
    if anchor is None:
        reasons.append("EVIDENCE_NOT_FOUND")
        if structural:
            reasons.append("DOCX_LOCATION_UNVERIFIED")
    if field != "amount_scope" and token is not None:
        if field.endswith("amount"):
            in_quote = any(canonical(field, m.group()) == token for m in re.finditer(r"(?<!\d)" + MONEY_TOKEN + r"(?![\d,])", quote))
        elif field == "business_registration_no":
            in_quote = any(canonical(field, m.group()) == token for m in re.finditer(r"\d{3}[\s-]*\d{2}[\s-]*\d{5}", quote))
        elif field == "issue_date":
            in_quote = any(canonical(field, m.group()) == token for m in re.finditer(r"\d{4}\s*[년./-]\s*\d{1,2}\s*[월./-]\s*\d{1,2}\s*일?", quote))
        else:
            in_quote = token in (canonical(field, quote) or "")
        if not in_quote:
            reasons.append("VALUE_NOT_IN_QUOTE")
    local = [item for item in options if item["page"] == page_no and item["line"] == line_no]
    matches = [item for item in local if item["value"] == token]
    if not matches:
        reasons.append("FIELD_LABEL_OR_VALUE_MISMATCH")
    if anchor and matches:
        anchor["label_line"] = matches[0]["label_line"]
        anchor["label"] = matches[0]["label"]
        anchor["party"] = matches[0]["party"]
    if field in {"company_name", "business_registration_no"} and matches and any(item["party"] != "supplier" for item in matches):
        reasons.append("WRONG_PARTY")
    if len(distinct) > 1:
        reasons.append("AMBIGUOUS")
    if field == "business_registration_no" and not checksum_valid(token):
        reasons.append("INVALID_CHECKSUM")
    if field == "amount_scope" and token not in {"FULL", "PARTIAL"}:
        reasons.append("INVALID_FORMAT")
    # Raw printed inconsistencies belong to the rule engine, not value repair.
    numbers = {key: {item["value"] for item in candidates(pages, key, document_type)} for key in ("supply_amount", "tax_amount", "total_amount")}
    if all(len(values) == 1 for values in numbers.values()):
        supply, tax, total = (Decimal(next(iter(numbers[key]))) for key in ("supply_amount", "tax_amount", "total_amount"))
        if supply + tax != total:
            warnings.append("SOURCE_ARITHMETIC_MISMATCH")
    return {"field": field, "proposal": proposal, "accepted": not reasons, "reasons": list(dict.fromkeys(reasons)),
            "warnings": warnings, "anchor": anchor, "human_confirmed": False}


def supplementation_targets(fields, pages, document_type):
    expected = {"company_name", "business_registration_no", "issue_date"}
    if document_type == "bank_account":
        expected.discard("business_registration_no")
    if document_type in {"contract", "estimate", "tax_invoice", "inspection_confirmation"}:
        expected.update(("supply_amount", "tax_amount", "total_amount", "amount_scope"))
    targets = set()
    for field in expected:
        current = canonical(field, fields.get(field))
        options = candidates(pages, field, document_type)
        eligible = [o for o in options if field not in {"company_name", "business_registration_no"} or o["party"] == "supplier"]
        choices = {o["value"] for o in eligible}
        if current is None or current not in choices or len(choices) > 1:
            targets.add(field)
    monetary = [canonical(key, fields.get(key)) for key in ("supply_amount", "tax_amount", "total_amount")]
    if all(value is not None for value in monetary) and Decimal(monetary[0]) + Decimal(monetary[1]) != Decimal(monetary[2]):
        targets.update(("supply_amount", "tax_amount", "total_amount"))
    return tuple(field for field in LABELS if field in targets)
