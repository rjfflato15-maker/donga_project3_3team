import os
from typing import Tuple, Any
from .input_safety import (
    UnsafeDocument, TextBudget, check_container, check_image_size, bounded_decompress, reject_xml_entities, checked_text,
    MAX_ARCHIVE_MEMBERS, MAX_ARCHIVE_EXPANDED_BYTES, MAX_EXTRACTED_TEXT_CHARS,
    MAX_SPREADSHEET_CELLS, MAX_SPREADSHEET_SHEETS,
)


class _BorrowedImageStream:
    """Allow Pillow to release its decoder without closing a caller's stream."""

    def __init__(self, stream):
        self.stream = stream

    def __getattr__(self, name):
        return getattr(self.stream, name)

    def close(self):
        pass


def _to_pil_image(image_input: Any):
    """Return an owned, oriented image; never close a caller's PIL image."""
    from PIL import Image, ImageOps
    import io

    original = None
    owned = not isinstance(image_input, Image.Image)
    image = None
    try:
        if isinstance(image_input, (str, os.PathLike)):
            original = Image.open(image_input)
        elif isinstance(image_input, (bytes, bytearray)):
            original = Image.open(io.BytesIO(image_input))
        elif isinstance(image_input, Image.Image):
            original = image_input
        elif hasattr(image_input, "read"):
            original = Image.open(_BorrowedImageStream(image_input))
        else:
            original = Image.fromarray(image_input)
        check_image_size(original)
        try:
            image = ImageOps.exif_transpose(original)
        except Exception:
            image = original.copy()
        image.load()
        # Transparent document backgrounds should render as white, like paper.
        if "A" in image.getbands() or "transparency" in image.info:
            with image.convert("RGBA") as rgba, Image.new("RGBA", image.size, "white") as background:
                with Image.alpha_composite(background, rgba) as composite:
                    converted = composite.convert("RGB")
            image.close()
            image = converted
        elif image.mode not in ("RGB", "L"):
            converted = image.convert("RGB")
            image.close()
            image = converted
        return image
    except Exception:
        if image is not None:
            image.close()
        return None
    finally:
        if owned and original is not None:
            original.close()


def _safe_winocr_recognize(target_img, lang: str = "ko"):
    """
    Safely executes WinOCR recognize_pil_sync.
    If called within a running asyncio event loop (e.g. FastAPI route),
    delegates execution to a worker thread to prevent:
    'RuntimeError: asyncio.run() cannot be called from a running event loop'
    """
    import winocr
    import asyncio
    from concurrent.futures import ThreadPoolExecutor

    try:
        loop = asyncio.get_running_loop()
    except RuntimeError:
        loop = None

    if loop is not None and loop.is_running():
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(winocr.recognize_pil_sync, target_img, lang)
            return future.result(timeout=45)
    else:
        return winocr.recognize_pil_sync(target_img, lang=lang)


def prepare_winocr_image(image):
    """Fit WinOCR's existing size budget without distorting aspect ratio.

    The result may be the input image. Callers own any new resized image.
    """
    from PIL import Image

    width, height = image.size
    if width <= 0 or height <= 0:
        raise ValueError("Invalid OCR image dimensions")
    scale = 1.0
    if (width < 1200 or height < 600) and max(width, height) * 2 <= 2500:
        scale = min(max(2, min(3, 1600 // width)), 2400 / max(width, height))
    elif max(width, height) > 2500:
        scale = 2400 / max(width, height)
    size = (max(1, round(width * scale)), max(1, round(height * scale)))
    return image if size == image.size else image.resize(size, Image.Resampling.LANCZOS)


def run_ocr_on_image(image_input: Any) -> str:
    """Run local Windows OCR, with the existing optional Tesseract fallback."""
    pil_img = _to_pil_image(image_input)
    if pil_img is None:
        return ""
    target_img = pil_img
    try:
        try:
            target_img = prepare_winocr_image(pil_img)
            result = _safe_winocr_recognize(target_img, lang="ko")
            if result:
                lines = [line.get("text", "").strip() for line in result.get("lines", [])
                         if line.get("text", "").strip()]
                if lines:
                    return "\n".join(lines).strip()
                if result.get("text", "").strip():
                    return result["text"].strip()
        except Exception:
            # Do not log exception details: OCR errors can contain source text.
            print("[Warning] WinOCR attempt failed")
        try:
            import pytesseract
            return pytesseract.image_to_string(pil_img, lang="kor+eng").strip()
        except Exception:
            return ""
    finally:
        if target_img is not pil_img:
            target_img.close()
        pil_img.close()


def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from Word .docx file (paragraphs and tables)."""
    check_container(file_path, ".docx")
    try:
        import docx
        doc = docx.Document(file_path)
        parts = TextBudget()
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                parts.append(t)
        for tbl in doc.tables:
            for row in tbl.rows:
                # Do not materialize another copy of long merged-cell text.
                deduped = []
                previous_cell = None
                for cell in row.cells:
                    if cell._tc is previous_cell:
                        continue
                    previous_cell = cell._tc
                    text = cell.text.strip()
                    if text and (not deduped or deduped[-1] != text):
                        deduped.append(text)
                if deduped:
                    parts.append("\t".join(deduped))
        return parts.join()
    except UnsafeDocument:
        raise
    except Exception as e:
        return ""


def extract_text_from_doc(file_path: str) -> str:
    """Extracts text from legacy Word .doc format using olefile and string decoding."""
    ole = None
    try:
        import olefile
        if olefile.isOleFile(file_path):
            ole = olefile.OleFileIO(file_path)
            if ole.exists("WordDocument"):
                raw = ole.openstream("WordDocument").read()
                u16 = raw.decode("utf-16le", errors="ignore")
                clean_lines = [l.strip() for l in u16.split("\n") if len(l.strip()) > 3]
                if clean_lines:
                    return "\n".join(clean_lines).strip()
    except UnsafeDocument:
        raise
    except Exception:
        pass
    finally:
        if ole is not None:
            ole.close()

    try:
        with open(file_path, "rb") as f:
            raw = f.read()
        u16 = raw.decode("utf-16le", errors="ignore")
        clean = [l.strip() for l in u16.split("\n") if len(l.strip()) > 3]
        if clean:
            return "\n".join(clean).strip()
        return raw.decode("utf-8", errors="ignore").strip()
    except UnsafeDocument:
        raise
    except Exception:
        return ""


def extract_text_from_hwpx(file_path: str) -> str:
    """Extracts text from Hancom .hwpx XML package."""
    import zipfile
    import xml.etree.ElementTree as ET
    import re

    text_parts = TextBudget()
    check_container(file_path, ".hwpx")
    try:
        with zipfile.ZipFile(file_path, "r") as zf:
            section_files = sorted(
                [f for f in zf.namelist() if f.startswith("Contents/section") and f.endswith(".xml")]
            )
            if not section_files:
                section_files = sorted(
                    [f for f in zf.namelist() if f.endswith(".xml") and not f.startswith("META-INF") and "version" not in f.lower()]
                )

            for sf in section_files:
                try:
                    xml_data = zf.read(sf)
                    reject_xml_entities(xml_data)
                    sec_text = ""
                    # 1. Try ElementTree
                    try:
                        root = ET.fromstring(xml_data)
                        paras = [elem for elem in root.iter() if elem.tag.endswith("}p") or elem.tag == "p"]
                        if paras:
                            lines = ["".join(p.itertext()).strip() for p in paras if "".join(p.itertext()).strip()]
                            sec_text = "\n".join(lines)
                        else:
                            sec_text = "".join(root.itertext()).strip()
                    except UnsafeDocument:
                        raise
                    except Exception:
                        pass

                    # 2. Regex fallback for unbound prefixes / namespaces
                    if not sec_text:
                        decoded = xml_data.decode("utf-8", errors="ignore")
                        t_matches = re.findall(r"<[^>]*:?t[^>]*>(.*?)</[^>]*:?t>", decoded, flags=re.DOTALL)
                        if t_matches:
                            lines = [re.sub(r"<[^>]+>", "", m).strip() for m in t_matches if m.strip()]
                            sec_text = "\n".join(lines)
                        else:
                            text_only = re.sub(r"<[^>]+>", " ", decoded)
                            lines = [l.strip() for l in text_only.split("\n") if l.strip()]
                            sec_text = "\n".join(lines)

                    if sec_text:
                        text_parts.append(sec_text)
                except UnsafeDocument:
                    raise
                except Exception:
                    continue
    except UnsafeDocument:
        raise
    except Exception:
        pass

    return text_parts.join("\n\n")


def extract_text_from_hwp(file_path: str) -> str:
    """Extracts text from Hancom .hwp binary format (HWP 5.0 OLE + zlib / pyhwp / string scanning)."""
    import zlib
    extracted_text = TextBudget()

    # 1. Try standard HWP 5.0 BodyText streams via olefile + zlib decompress
    ole = None
    try:
        import olefile
        if olefile.isOleFile(file_path):
            ole = olefile.OleFileIO(file_path)
            section_names = []
            for stream in ole.listdir():
                if len(stream) >= 2 and stream[0] == "BodyText" and stream[1].startswith("Section"):
                    section_names.append("/".join(stream))
            section_names.sort()
            if len(section_names) > MAX_ARCHIVE_MEMBERS:
                raise UnsafeDocument('HWP 구역 수가 처리 한도를 초과했습니다.')
            expanded_bytes = 0

            for sec_name in section_names:
                try:
                    stream_data = ole.openstream(sec_name).read()
                    try:
                        decompressed = bounded_decompress(stream_data, -15)
                    except UnsafeDocument:
                        raise
                    except Exception:
                        try:
                            decompressed = bounded_decompress(stream_data, zlib.MAX_WBITS)
                        except UnsafeDocument:
                            raise
                        except Exception:
                            decompressed = stream_data

                    expanded_bytes += len(decompressed)
                    if expanded_bytes > MAX_ARCHIVE_EXPANDED_BYTES:
                        raise UnsafeDocument('HWP 문서의 펼친 크기가 처리 한도를 초과했습니다.')
                    offset = 0
                    sec_text = []
                    while offset < len(decompressed) - 4:
                        header = int.from_bytes(decompressed[offset:offset+4], "little")
                        tag_id = header & 0x3FF
                        size = (header >> 20) & 0xFFF
                        offset += 4
                        if size == 0xFFF:
                            if offset + 4 > len(decompressed):
                                break
                            size = int.from_bytes(decompressed[offset:offset+4], "little")
                            offset += 4
                        if offset + size > len(decompressed):
                            break
                        payload = decompressed[offset:offset+size]
                        offset += size

                        if tag_id == 67:  # HWPTAG_PARA_TEXT
                            chars = []
                            for i in range(0, len(payload) - 1, 2):
                                code = int.from_bytes(payload[i:i+2], "little")
                                if code in (10, 13):
                                    chars.append("\n")
                                elif code >= 32:
                                    chars.append(chr(code))
                            line = "".join(chars).strip()
                            if line:
                                sec_text.append(line)

                    if sec_text:
                        extracted_text.append("\n".join(sec_text))
                except UnsafeDocument:
                    raise
                except Exception:
                    continue
    except UnsafeDocument:
        raise
    except Exception:
        pass
    finally:
        if ole is not None:
            ole.close()

    # 2. Try pyhwp cli tool as secondary
    if not extracted_text.parts:
        try:
            import subprocess
            res = subprocess.run(["hwp5txt", file_path], capture_output=True, text=True, timeout=5)
            if res.returncode == 0 and res.stdout.strip():
                return res.stdout.strip()
        except UnsafeDocument:
            raise
        except Exception:
            pass

    # 3. String scanning fallback for UTF-16LE text
    if not extracted_text.parts:
        try:
            with open(file_path, "rb") as f:
                raw_bytes = f.read()
            u16 = raw_bytes.decode("utf-16le", errors="ignore")
            clean_lines = [line.strip() for line in u16.split("\n") if len(line.strip()) > 3]
            if clean_lines:
                extracted_text.append("\n".join(clean_lines))
        except UnsafeDocument:
            raise
        except Exception:
            pass

    return extracted_text.join("\n\n")


def extract_text_from_xlsx(file_path: str) -> str:
    """Extract spreadsheets without trusting attacker-controlled dimensions."""
    workbook = None
    try:
        import openpyxl
        check_container(file_path, ".xlsx")
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True, keep_links=False)
        if len(workbook.sheetnames) > MAX_SPREADSHEET_SHEETS:
            raise UnsafeDocument("시트 수가 처리 한도를 초과했습니다.")
        output = TextBudget()
        scanned_cells = 0
        for sheet in workbook:
            rows, columns = sheet.max_row, sheet.max_column
            if (rows is not None and columns is not None
                    and rows * columns > MAX_SPREADSHEET_CELLS - scanned_cells):
                raise UnsafeDocument("스프레드시트 범위가 처리 한도를 초과했습니다. 사용하지 않는 행·열을 정리하거나 문서를 나누세요.")
            lines = TextBudget()
            for row in sheet.iter_rows(values_only=True):
                scanned_cells += max(len(row), 1)
                if scanned_cells > MAX_SPREADSHEET_CELLS:
                    raise UnsafeDocument("스프레드시트 셀 수가 처리 한도를 초과했습니다.")
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    lines.append("\t".join(cells))
            if lines.parts:
                output.append(f"[시트: {sheet.title}]\n" + lines.join())
        return output.join("\n\n")
    except UnsafeDocument:
        raise
    except Exception:
        return ""
    finally:
        if workbook is not None:
            workbook.close()


def extract_text_from_xls(file_path: str) -> str:
    """Extracts text from legacy Excel .xls spreadsheet."""
    try:
        import xlrd
        wb = xlrd.open_workbook(file_path)
        if wb.nsheets > MAX_SPREADSHEET_SHEETS:
            raise UnsafeDocument('시트 수가 처리 한도를 초과했습니다.')
        if sum(sheet.nrows * sheet.ncols for sheet in wb.sheets()) > MAX_SPREADSHEET_CELLS:
            raise UnsafeDocument('스프레드시트 셀 수가 처리 한도를 초과했습니다.')
        sheet_texts = TextBudget()
        for sheet in wb.sheets():
            row_lines = []
            for row_idx in range(sheet.nrows):
                row_vals = sheet.row_values(row_idx)
                cells = [str(val).strip() for val in row_vals if val is not None and str(val).strip() != ""]
                if cells:
                    row_lines.append("\t".join(cells))
            if row_lines:
                sheet_texts.append(f"[시트: {sheet.name}]\n" + "\n".join(row_lines))
        return sheet_texts.join("\n\n")
    except UnsafeDocument:
        raise
    except Exception:
        return ""


def extract_text_from_csv(file_path: str) -> str:
    """Extracts text from .csv spreadsheet with multi-encoding detection."""
    import csv
    encodings = ["utf-8-sig", "cp949", "euc-kr", "utf-8"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                lines = TextBudget()
                for row in reader:
                    cells = [c.strip() for c in row if c.strip()]
                    if cells:
                        lines.append("\t".join(cells))
                if lines.parts:
                    return lines.join()
        except UnsafeDocument:
            raise
        except Exception:
            continue
    return ""


MAX_PDF_OCR_PAGES = 50
MAX_PDF_NATIVE_PAGES = 500
MAX_PDF_TEXT_CHARS = 2_000_000
MAX_PDF_OCR_EDGE = 2400


def _pdf_page_record(number, text=""):
    return {"page": number, "location_kind": "page", "text": text,
            "evidence_complete": False, "verified_blank": False, "extraction_warnings": []}


def _pdf_warning(page, code):
    page["evidence_complete"] = False
    if code not in page["extraction_warnings"]:
        page["extraction_warnings"].append(code)


def _merge_pdf_text(native, recognized):
    """Retain complete native and OCR blocks without rewriting their order.

    Removing repeated lines can erase party boundaries or field labels. Only
    deduplicate when the entire trimmed blocks match exactly.
    """
    native, recognized = native.strip(), recognized.strip()
    if not native:
        return recognized
    if not recognized or native == recognized:
        return native
    return native + "\n\n" + recognized


def _pdf_ocr_scale(width, height):
    """Render at the useful OCR size directly, with bounded pixel allocation."""
    import math
    if not all(math.isfinite(value) and value > 0 for value in (width, height)):
        raise ValueError("Invalid PDF page dimensions")
    scale = min(2.0, MAX_PDF_OCR_EDGE / max(width, height))
    if not math.isfinite(scale) or scale <= 0:
        raise ValueError("Invalid PDF render scale")
    return scale


def extract_pdf_pages(file_path: str) -> list[dict[str, Any]]:
    """Extract physical pages once per pass, retaining incomplete evidence.

    A native footer cannot suppress OCR of an image, Form, outline or
    annotation. Native-only pages avoid OCR. The renderer is opened once and
    each page is rendered at most once, outside any OCR engine call/lock.
    OCR is bounded to documents of at most 50 pages; native text remains
    available up to 500 pages / 2M characters. No raw-text cache is retained.
    Completeness means extraction ran, not that recognition is error-free.
    """
    pages = []
    try:
        from pypdf import PdfReader
        with open(file_path, "rb") as source:
            reader = PdfReader(source)
            count = len(reader.pages)
            if count > MAX_PDF_NATIVE_PAGES:
                return []
            pages = [_pdf_page_record(index + 1) for index in range(count)]
            remaining = MAX_PDF_TEXT_CHARS
            for index, record in enumerate(pages):
                try:
                    native = (reader.pages[index].extract_text() or "").strip()
                    record["text"] = native[:remaining]
                    remaining -= len(record["text"])
                    if len(native) > len(record["text"]):
                        _pdf_warning(record, "PDF_TEXT_LIMIT")
                except Exception:
                    # PDFium may still recover this page's native/raster text.
                    pass
    except Exception:
        pass

    try:
        from .pdf_runtime import (open_pdf, open_page, open_textpage, page_count,
                                  page_size, page_text, page_is_blank, page_requires_raster_masking,
                                  render_page_image)
        with open_pdf(file_path) as pdf:
            count = page_count(pdf)
            if count > MAX_PDF_NATIVE_PAGES:
                return []
            mismatch = bool(pages) and len(pages) != count
            while len(pages) < count:
                pages.append(_pdf_page_record(len(pages) + 1))
            if max(count, len(pages)) > MAX_PDF_OCR_PAGES:
                for record in pages:
                    _pdf_warning(record, "PDF_PAGE_LIMIT")
                return pages
            for index in range(count):
                record = pages[index]
                image = None
                try:
                    with open_page(pdf, index) as page:
                        if not record["text"] and "PDF_TEXT_LIMIT" not in record["extraction_warnings"]:
                            try:
                                with open_textpage(page) as textpage:
                                    record["text"] = page_text(textpage).strip()
                            except Exception:
                                pass
                        record["verified_blank"] = not record["text"].strip() and page_is_blank(page)
                        needs_ocr = not record["verified_blank"] and (
                            not record["text"].strip() or page_requires_raster_masking(page))
                        if needs_ocr:
                            scale = _pdf_ocr_scale(*page_size(page))
                            image = render_page_image(page, scale)
                        else:
                            image = None
                    # Close the native page handle before running the OCR engine.
                    if image is not None:
                        recognized = run_ocr_on_image(image)
                        if recognized.strip():
                            record["text"] = _merge_pdf_text(record["text"], recognized)
                        else:
                            _pdf_warning(record, "PDF_OCR_EMPTY")
                    record["evidence_complete"] = not record["extraction_warnings"]
                except Exception:
                    _pdf_warning(record, "PDF_PAGE_EXTRACTION_FAILED")
                finally:
                    if image is not None:
                        image.close()
            if mismatch:
                for record in pages:
                    _pdf_warning(record, "PDF_PAGE_COUNT_MISMATCH")
    except Exception:
        # Preserve native text if rendering is unavailable, but never claim
        # that uninspected raster content was absent from those pages.
        for record in pages:
            _pdf_warning(record, "PDF_RENDERER_UNAVAILABLE")
        if not pages or any(not record["text"] for record in pages):
            try:
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    count = len(pdf.pages)
                    if count > MAX_PDF_NATIVE_PAGES:
                        return []
                    while len(pages) < count:
                        pages.append(_pdf_page_record(len(pages) + 1))
                    for index, page in enumerate(pdf.pages):
                        record = pages[index]
                        _pdf_warning(record, "PDF_RENDERER_UNAVAILABLE")
                        if not record["text"]:
                            try:
                                record["text"] = (page.extract_text() or "").strip()
                            except Exception:
                                pass
            except Exception:
                pass
    remaining = MAX_PDF_TEXT_CHARS
    for record in pages:
        text = record["text"]
        record["text"] = text[:remaining]
        remaining -= len(record["text"])
        if len(text) > len(record["text"]):
            _pdf_warning(record, "PDF_TEXT_LIMIT")
    return pages


def extract_text_from_file(file_path: str, file_name: str = "", *, extraction_metadata: dict | None = None) -> Tuple[str, str]:
    """
    Extracts text from various file formats:
    - PDF (.pdf)
    - Images (.png, .jpg, .jpeg, .bmp, .webp, .tiff, .tif, .gif)
    - Word (.docx, .doc)
    - Hancom (.hwp, .hwpx)
    - Excel (.xlsx, .xls, .csv)
    - Plain text (.txt, .md, .json)
    Returns (extracted_text, file_type)
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    if not file_name:
        file_name = os.path.basename(file_path)
    # A multipart filename is display metadata, never extra source-text lines.
    from ..services.storage_service import clean_name
    file_name = clean_name(file_name)

    ext = os.path.splitext(file_path)[1].lower()
    check_container(file_path, ext)
    header = f"[파일명: {file_name}]\n"

    # 1. PDF Files. Keep every page's text, not just the first nonempty layer.
    if ext == ".pdf":
        pages = extract_pdf_pages(file_path)
        pdf_text = "\n\n".join(page["text"] for page in pages if page["text"]).strip()
        if extraction_metadata is not None:
            incomplete = [page for page in pages if not (
                page["evidence_complete"] and (page["text"].strip() or page.get("verified_blank") is True))]
            warnings = list(dict.fromkeys(code for page in pages for code in page["extraction_warnings"]))
            if not pages:
                warnings.append("PDF_SOURCE_UNAVAILABLE")
            elif not pdf_text:
                warnings.append("PDF_NO_READABLE_TEXT")
            extraction_metadata.update(
                evidence_complete=bool(pdf_text) and not incomplete,
                page_count=len(pages),
                incomplete_pages=[page["page"] for page in incomplete][:MAX_PDF_NATIVE_PAGES],
                warning_codes=warnings[:32],
            )
        return (header + pdf_text).strip(), "pdf"

    # 2. Image files (.png, .jpg, .jpeg, .bmp, .webp, .tiff, .tif, .gif)
    if ext in [".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tiff", ".tif", ".gif"]:
        ocr_text = run_ocr_on_image(file_path)
        full_text = header + checked_text(ocr_text or "")
        return full_text.strip(), ext.replace(".", "")

    # 3. Word files (.docx, .doc)
    if ext == ".docx":
        docx_text = extract_text_from_docx(file_path)
        return (header + checked_text(docx_text)).strip(), "docx"
    if ext == ".doc":
        doc_text = extract_text_from_doc(file_path)
        return (header + checked_text(doc_text)).strip(), "doc"

    # 4. Hancom files (.hwpx, .hwp)
    if ext == ".hwpx":
        hwpx_text = extract_text_from_hwpx(file_path)
        return (header + checked_text(hwpx_text)).strip(), "hwpx"
    if ext == ".hwp":
        hwp_text = extract_text_from_hwp(file_path)
        return (header + checked_text(hwp_text)).strip(), "hwp"

    # 5. Excel files (.xlsx, .xlsm, .xls, .csv)
    if ext in [".xlsx", ".xlsm"]:
        xlsx_text = extract_text_from_xlsx(file_path)
        return (header + checked_text(xlsx_text)).strip(), "xlsx"
    if ext == ".xls":
        xls_text = extract_text_from_xls(file_path)
        return (header + checked_text(xls_text)).strip(), "xls"
    if ext == ".csv":
        csv_text = extract_text_from_csv(file_path)
        return (header + checked_text(csv_text)).strip(), "csv"

    # 6. Default text reading with encoding fallback
    encodings = ["utf-8-sig", "utf-8", "cp949", "euc-kr", "latin-1"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read(MAX_EXTRACTED_TEXT_CHARS + 1)
                if len(content) > MAX_EXTRACTED_TEXT_CHARS:
                    raise UnsafeDocument('문서 텍스트가 처리 한도를 초과했습니다. 문서를 나눠 등록하세요.')
                return (header + content).strip(), ext.replace(".", "") or "txt"
        except UnsafeDocument:
            raise
        except Exception:
            continue

    # 7. Binary fallback
    with open(file_path, "rb") as f:
        raw = f.read()
        return (header + raw.decode("utf-8", errors="ignore")).strip(), "bin"

