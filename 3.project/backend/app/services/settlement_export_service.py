"""Saved settlement values and their review state; no ERP posting or new AI calls."""
import hashlib
import io
import json
from datetime import datetime, timedelta, timezone

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..db.models import Contract
from ..schemas.document import AnalysisFields
from .contract_service import ContractService
from .contract_export_service import collect_sensitive, public_data
from .document_version_service import active_documents, metadata
from .field_confirmation_service import current_confirmations
from .storage_service import STORAGE_LOCK
from .workflow_service import audit

SEOUL = timezone(timedelta(hours=9))
DOCUMENT_NAMES = {'contract': '계약서', 'estimate': '견적서',
    'business_registration': '사업자등록증', 'bank_account': '통장사본',
    'tax_invoice': '세금계산서', 'inspection_confirmation': '검수확인서', 'unknown': '기타 서류'}
FIELD_NAMES = {'company_name': '상호', 'business_registration_no': '사업자번호',
    'supply_amount': '공급가액', 'tax_amount': '세액', 'total_amount': '합계',
    'bank_name': '은행', 'account_number': '계좌번호', 'issue_date': '발행일',
    'amount_scope': '청구 범위'}
STATUS_NAMES = {'PASS': '적용 검사 일치', 'REVIEW': '확인 필요', 'FAIL': '불일치',
    'MISSING': '누락', 'PENDING': '확인 전', 'CONFIRMED': '최종 확인 완료',
    'RETURNED': '보완 요청', 'STALE': '자료 변경 · 재확인',
    'RECORDED': '지급 기록 완료', 'RECONFIRMATION_REQUIRED': '기존 지급 재확인 필요', 'VOIDED': '지급 기록 취소'}


def seoul_time(value):
    if not value:
        return None
    moment = datetime.fromisoformat(value) if isinstance(value, str) else value
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(SEOUL).strftime('%Y-%m-%d %H:%M:%S')


def capture(db, cid, revision):
    detail = ContractService().get_contract_detail(db, cid)
    if detail is None:
        raise HTTPException(404, '계약을 찾을 수 없습니다.')
    data = detail.model_dump(mode='json')
    workflow = data['workflow']
    if revision is not None and workflow['revision'] != revision:
        raise HTTPException(409, '계약 자료가 변경되었습니다. 최신 자료를 다시 불러오세요.')
    documents = []
    for doc in sorted(active_documents(db.get(Contract, cid)), key=lambda item: item.document_id):
        stored = json.loads(doc.extract.raw_fields or '{}') if doc.extract else {}
        documents.append({'document_id': doc.document_id, 'file_name': doc.file_name,
            'document_type': doc.document_type, 'processing_status': doc.processing_status,
            'version': metadata(doc), 'fields': AnalysisFields(**stored).model_dump(),
            'confirmations': current_confirmations(db, doc.document_id, stored),
            'ai_review': stored.get('ai_review') or {}})
    if len(documents) > 100:
        raise HTTPException(413, '정산 엑셀은 계약 한 건의 현재 서류 100개까지 지원합니다.')
    snapshot = {'contract': data, 'documents': documents}
    snapshot['signature'] = hashlib.sha256(json.dumps(snapshot, sort_keys=True,
        ensure_ascii=False, default=str).encode()).hexdigest()
    db.rollback()
    return snapshot


def sheets_for(snapshot, masked):
    c = snapshot['contract']
    docs = snapshot['documents']
    wf = c['workflow']
    invoices = [d for d in docs if d['document_type'] == 'tax_invoice']
    accounts = [d for d in docs if d['document_type'] == 'bank_account']
    # Independent invoices/accounts are not replacements. Never silently choose one.
    invoice = invoices[0]['fields'] if len(invoices) == 1 else {}
    account = accounts[0]['fields'] if len(accounts) == 1 else {}
    reviews = wf.get('reviews') or []
    current_review = next((r for r in reviews if r['revision'] == wf['revision']), None)
    approved_on = seoul_time(current_review['created_at']) if (
        wf['approval_status'] == 'CONFIRMED' and current_review and current_review['decision'] == 'CONFIRMED') else None
    payment = c.get('payment') or {}
    notes = []
    if len(invoices) != 1:
        notes.append('세금계산서 없음' if not invoices else '세금계산서 여러 건: 서류별 값에서 확인, 합산하지 않음')
    if len(accounts) != 1:
        notes.append('통장사본 없음' if not accounts else '통장사본 여러 건: 서류별 값에서 지급 계좌 확인')
    if invoice.get('amount_scope') != 'FULL':
        notes.append('세금계산서의 전액·부분 청구 범위 확인 필요')
    summary = [c['contract_id'], c['title'], c['vendor_name'], c['business_number'],
        c['contract_amount'], {'TOTAL': '부가세 포함 합계', 'SUPPLY': '공급가액', 'UNKNOWN': '미확인'}[wf['amount_basis']],
        invoice.get('supply_amount'), invoice.get('tax_amount'), invoice.get('total_amount'),
        account.get('bank_name'), account.get('account_number'), STATUS_NAMES.get(c['review_status'], c['review_status']),
        STATUS_NAMES.get(wf['approval_status'], wf['approval_status']), approved_on,
        STATUS_NAMES.get(payment.get('status'), '지급 기록 없음'), payment.get('paid_on'),
        payment.get('recorded_paid_on'), wf['revision'], ' / '.join(notes)]
    doc_rows, field_rows = [], []
    for d in docs:
        f = d['fields']
        proofs = d['confirmations']
        doc_rows.append([d['document_id'], DOCUMENT_NAMES.get(d['document_type'], d['document_type']),
            d['file_name'], d['version']['version_number'], d['processing_status'],
            f.get('company_name'), f.get('business_registration_no'), f.get('issue_date'),
            f.get('supply_amount'), f.get('tax_amount'), f.get('total_amount'),
            {'FULL': '전액 청구', 'PARTIAL': '부분 청구'}.get(f.get('amount_scope'), '미확인'),
            f.get('bank_name'), f.get('account_number')])
        decisions = {x.get('field'): x for x in d['ai_review'].get('decisions', []) if isinstance(x, dict)}
        for key, label in FIELD_NAMES.items():
            value = f.get(key)
            proof = proofs.get(key) or {}
            decision = decisions.get(key) or {}
            source = '사람 확인' if proof else 'AI 보완' if decision.get('applied') else '규칙 추출'
            if value in (None, '', 'UNKNOWN'):
                value, source = None, '미추출·미확인'
            field_rows.append([d['document_id'], DOCUMENT_NAMES.get(d['document_type'], d['document_type']),
                label, value, source, '확인 완료' if proof else '미확인',
                proof.get('actor_name'), seoul_time(proof.get('at'))])
    tables = {
        '정산 내역': (["계약 ID", "계약명", "등록 거래처", "등록 사업자번호", "계약 금액", "금액 기준",
            "세금계산서 공급가액", "세금계산서 세액", "세금계산서 합계", "은행", "계좌번호",
            "자동 검사 상태", "대표 최종 확인 상태", "현재 자료 최종 확인일 (한국시간)",
            "지급 기록 상태", "현재 자료 지급일", "기존 지급 기록일", "자료 버전", "확인 사항"], [summary]),
        '서류별 값': (["서류 ID", "서류 종류", "파일명", "서류 버전", "처리 상태", "서류 상호",
            "서류 사업자번호", "발행일", "공급가액", "세액", "합계", "청구 범위", "은행", "계좌번호"], doc_rows),
        '필드 확인': (["서류 ID", "서류 종류", "항목", "저장값", "값 출처", "사람 확인", "확인자", "확인일 (한국시간)"], field_rows),
        '자동 검사': (["검사 ID", "상태", "관련 서류·항목", "기준값", "실제값", "검사 설명"],
            [[x['rule_id'], STATUS_NAMES.get(x['status'], x['status']), x['field'], x.get('expected'), x.get('actual'), x['message']] for x in c['checks']]),
        '안내': (["항목", "내용"], [
            ['생성 시각 (한국시간)', datetime.now(SEOUL).strftime('%Y-%m-%d %H:%M:%S')],
            ['자료 범위', '이 계약의 현재 서류와 저장값·검사·결재·지급 상태. 이전 본은 증빙 ZIP에서 확인하세요.'],
            ['미확인 값', '빈 칸은 미추출·미확인 또는 여러 서류 중 선택하지 않은 값입니다. 0원과 구분합니다.'],
            ['금액', '문서에 저장된 공급가액·세액·합계만 내보냅니다. 세액 역산이나 여러 계산서 합산을 하지 않습니다.'],
            ['검수와 결재', '자동 검사·필드별 사람 확인·대표 최종 확인은 별도입니다. 결재가 자동 검사 상태를 바꾸지 않습니다.'],
            ['지급 기록', '실제 송금 기능이 아닙니다. 기존 지급 기록일이 있으면 새 지급 전 기존 기록을 확인하세요.'],
            ['회계 전달', '일반 정산 자료입니다. 특정 ERP의 전표 업로드 양식이 아니며 계정과목·거래처 코드 등의 매핑이 필요합니다.'],
            ['가림', '가림본: 회계 입력에는 실제값이 필요합니다.' if masked else '실제값 포함: 업무에 필요한 대상에게 전달하세요.'],
        ]),
    }
    if masked:
        sensitive = {}
        collect_sensitive(snapshot, sensitive)
        return public_data({name: list(table) for name, table in tables.items()}, True, sensitive)
    return tables


def workbook_bytes(tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (headers, rows) in tables.items():
        sheet = workbook.create_sheet(name)
        for row_index, values in enumerate([headers, *rows], 1):
            for col_index, value in enumerate(values, 1):
                cell = sheet.cell(row_index, col_index)
                # Explicit strings preserve leading zeros and cannot become formulas.
                if isinstance(value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub('', value)[:32767]
                    cell.data_type = 's'
                    cell.number_format = '@'
                else:
                    cell.value = value
                    if isinstance(value, (int, float)):
                        cell.number_format = '#,##0.########'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_index == 1:
                    cell.font = Font(name='Malgun Gothic', bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='334155')
                else:
                    cell.font = Font(name='Malgun Gothic', size=11)
            sheet.row_dimensions[row_index].height = 42 if row_index == 1 else 36
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 46 if header in {'내용', '확인 사항', '검사 설명', '파일명', '계약명'} else 23
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def create_settlement(db, cid, actor, masked=True, revision=None):
    with STORAGE_LOCK:
        snapshot = capture(db, cid, revision)
    content = workbook_bytes(sheets_for(snapshot, masked))
    with STORAGE_LOCK:
        db.expire_all()
        latest = capture(db, cid, revision)
        if latest['signature'] != snapshot['signature']:
            raise HTTPException(409, '엑셀 생성 중 계약 자료가 변경되었습니다. 다시 불러오세요.')
        audit(db, actor, 'SETTLEMENT_EXPORT_REQUEST', cid,
            details={'masked': masked, 'revision': snapshot['contract']['workflow']['revision'],
                     'document_count': len(snapshot['documents'])})
    return content
