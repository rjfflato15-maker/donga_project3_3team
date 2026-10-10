"""Small synthetic regressions; never allocate the hostile declared dimensions."""
import io
import json
import re
import zipfile
import zlib
from pathlib import Path

import openpyxl
import pytest
from fastapi import HTTPException, UploadFile
from PIL import Image

from backend.app.ai import input_safety, preprocessor
from backend.app.services import storage_service


def package(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, contents in entries.items():
            entry = zipfile.ZipInfo(name)
            # ZipInfo normalizes Windows separators; retain the literal test
            # input so malformed archive paths are exercised on every OS.
            entry.filename = name
            entry.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(entry, contents)
    return output.getvalue()


def sheet_package(rewrite):
    workbook = openpyxl.Workbook()
    workbook.active['A1'] = 'small harmless fixture'
    original = io.BytesIO()
    workbook.save(original)
    output = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            contents = source.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                contents = rewrite(contents)
            target.writestr(entry, contents)
    return output.getvalue()


@pytest.mark.parametrize('name,content', [
    ('fake.png', b'not an image'), ('fake.docx', b'PKgarbage'),
    ('fake.pdf', b'not PDF'), ('fake.hwp', b'plain text'),
    ('fake.doc', b'plain text'), ('fake.xls', b'plain text'),
    ('wrong.xlsx', package({'Contents/section0.xml': '<p>HWPX text</p>'})),
])
def test_format_confusion_is_rejected(name, content):
    with pytest.raises(HTTPException) as error:
        storage_service.validate_upload(name, content)
    assert error.value.status_code == 400


def test_image_extension_must_match_header():
    content = io.BytesIO()
    Image.new('RGB', (2, 2)).save(content, 'PNG')
    with pytest.raises(HTTPException):
        storage_service.validate_upload('mislabeled.jpg', content.getvalue())
    storage_service.validate_upload('real.png', content.getvalue())


@pytest.mark.parametrize('limit', ['MAX_ARCHIVE_MEMBER_BYTES', 'MAX_ARCHIVE_EXPANDED_BYTES'])
def test_compressed_office_limit_is_checked_before_reading_member(monkeypatch, limit):
    payload = package({'Contents/section0.xml': '<p>' + 'A' * 4096 + '</p>'})
    assert len(payload) < 1024
    monkeypatch.setattr(input_safety, limit, 1024)
    monkeypatch.setattr(zipfile.ZipFile, 'read', lambda *_a, **_k: pytest.fail('expanded before safety check'))
    with pytest.raises(HTTPException):
        storage_service.validate_upload('compressed.hwpx', payload)


def test_archive_member_count_has_preflight_budget(monkeypatch):
    payload = package({f'Contents/section{i}.xml': '<p>tiny</p>' for i in range(3)})
    monkeypatch.setattr(input_safety, 'MAX_ARCHIVE_MEMBERS', 2)
    with pytest.raises(HTTPException):
        storage_service.validate_upload('many.hwpx', payload)


@pytest.mark.parametrize('entry', ['../escape.xml', '/escape.xml', 'Contents/../escape.xml', 'Contents\\escape.xml'])
def test_office_archive_paths_are_rejected(entry):
    payload = package({entry: '<p>small</p>'})
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.infolist()[0].orig_filename == entry
    with pytest.raises(HTTPException):
        storage_service.validate_upload('paths.hwpx', payload)


@pytest.mark.parametrize('rewrite', [
    lambda data: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', data),
    lambda data: re.sub(rb'<dimension[^>]*/>', b'', data).replace(b'r="A1"', b'r="XFD1048576"'),
    lambda data: re.sub(rb'<dimension[^>]*/>', b'', data).replace(b'r="A1"', b'r="AAAAAA1"'),
])
def test_sparse_spreadsheet_rejected_before_openpyxl_iteration(rewrite, tmp_path, monkeypatch):
    payload = sheet_package(rewrite)
    assert len(payload) < 10_000
    path = tmp_path / 'tiny-hostile.xlsx'
    path.write_bytes(payload)
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *_a, **_k: pytest.fail('spreadsheet decoder was reached'))
    with pytest.raises(input_safety.UnsafeDocument):
        preprocessor.extract_text_from_xlsx(str(path))
    with pytest.raises(HTTPException):
        storage_service.validate_upload(path.name, payload)


def test_small_workbook_is_still_extracted(tmp_path):
    path = tmp_path / 'normal.xlsx'
    path.write_bytes(sheet_package(lambda data: data))
    storage_service.validate_upload(path.name, path.read_bytes())
    assert 'small harmless fixture' in preprocessor.extract_text_from_xlsx(str(path))


def test_xml_entities_rejected_before_workbook_decode():
    payload = sheet_package(lambda data: b'<!DOCTYPE worksheet [<!ENTITY example "text">]>' + data)
    with pytest.raises(HTTPException):
        storage_service.validate_upload('entities.xlsx', payload)


def test_image_pixel_limit_precedes_decode(monkeypatch):
    content = io.BytesIO()
    Image.new('RGB', (20, 20)).save(content, 'PNG')
    monkeypatch.setattr(input_safety, 'MAX_IMAGE_PIXELS', 100)
    monkeypatch.setattr(Image.Image, 'load', lambda *_a, **_k: pytest.fail('pixel decode was reached'))
    with pytest.raises(HTTPException):
        storage_service.validate_upload('pixels.png', content.getvalue())


def test_hwp_decompression_has_output_budget(monkeypatch):
    monkeypatch.setattr(input_safety, 'MAX_ARCHIVE_MEMBER_BYTES', 1024)
    compressed = zlib.compress(b'x' * 4096)
    with pytest.raises(input_safety.UnsafeDocument):
        input_safety.bounded_decompress(compressed, zlib.MAX_WBITS)
    assert input_safety.bounded_decompress(zlib.compress(b'normal'), zlib.MAX_WBITS) == b'normal'


def test_hwp_limit_is_not_swallowed_by_fallback(tmp_path, monkeypatch):
    import olefile
    import subprocess
    class TinyOle:
        def listdir(self): return [['BodyText', 'Section0']]
        def openstream(self, _): return io.BytesIO(zlib.compress(b'x' * 4096))
        def close(self): pass
    path = tmp_path / 'tiny.hwp'
    path.write_bytes(b'tiny synthetic placeholder')
    monkeypatch.setattr(olefile, 'isOleFile', lambda _: True)
    monkeypatch.setattr(olefile, 'OleFileIO', lambda _: TinyOle())
    monkeypatch.setattr(input_safety, 'MAX_ARCHIVE_MEMBER_BYTES', 1024)
    monkeypatch.setattr(subprocess, 'run', lambda *_a, **_k: pytest.fail('unsafe fallback was invoked'))
    with pytest.raises(input_safety.UnsafeDocument):
        preprocessor.extract_text_from_hwp(str(path))


def test_plain_text_limit_is_not_swallowed_by_encoding_fallback(tmp_path, monkeypatch):
    path = tmp_path / 'long.txt'
    path.write_text('x' * 65)
    monkeypatch.setattr(preprocessor, 'MAX_EXTRACTED_TEXT_CHARS', 64)
    with pytest.raises(input_safety.UnsafeDocument):
        preprocessor.extract_text_from_file(str(path))


def test_file_read_and_filename_storage_remain_bounded(monkeypatch):
    monkeypatch.setattr(storage_service, 'MAX_UPLOAD_BYTES', 16)
    file = UploadFile(filename='a.txt', file=io.BytesIO(b'x' * 100))
    with pytest.raises(HTTPException) as error:
        storage_service.read_upload(file)
    assert error.value.status_code == 413
    assert file.file.tell() == 17
    assert storage_service.clean_name('../../escape.txt') == 'escape.txt'
    assert storage_service.clean_name(r'C:\outside\escape.txt') == 'escape.txt'
    name = storage_service.clean_name('a' * 300 + '.pdf')
    assert len(name) == 240 and Path(name).suffix == '.pdf'


def test_bad_upload_has_no_stored_original(operator_client, db_factory):
    from test_api import create_contract
    from backend.app.db.models import Document
    contract = create_contract(operator_client)
    response = operator_client.post(f"/api/contracts/{contract['contract_id']}/documents",
                                   files={'files': ('broken.docx', b'PKgarbage', 'application/octet-stream')})
    assert response.status_code == 400
    with db_factory() as db:
        assert db.query(Document).count() == 0
    assert not list(storage_service.STORAGE_DIR.rglob('*.docx'))


@pytest.mark.parametrize('entry', ['raw/NUL.txt', 'raw/CON', 'raw./file.txt', 'raw/space /file.txt',
                                   'raw//same.txt', 'raw/./same.txt', 'raw/a:stream'])
def test_restore_windows_aliases_rejected_before_writes(entry, tmp_path):
    import manage_pilot as pilot
    archive = tmp_path / 'bad-backup.zip'
    archive.write_bytes(package({'pilot.db': 'placeholder', 'manifest.json': '{}', entry: 'tiny'}))
    destination = tmp_path / 'new-inbox'
    with pytest.raises(RuntimeError):
        pilot.restore(archive, destination)
    assert not destination.exists()


@pytest.mark.parametrize('limit', ['MAX_BACKUP_MANIFEST_BYTES', 'MAX_BACKUP_EXPANDED_BYTES'])
def test_restore_limits_precede_manifest_decompression(limit, tmp_path, monkeypatch):
    import manage_pilot as pilot
    archive = tmp_path / 'oversized-manifest.zip'
    manifest = json.dumps({'format': 'c-pilot-v1', 'database': 'pilot.db', 'padding': 'x' * 2048})
    archive.write_bytes(package({'pilot.db': 'placeholder', 'manifest.json': manifest}))
    monkeypatch.setattr(pilot, limit, 1024)
    monkeypatch.setattr(zipfile.ZipFile, 'read', lambda *_a, **_k: pytest.fail('manifest expanded before size check'))
    destination = tmp_path / 'new-inbox'
    with pytest.raises(RuntimeError):
        pilot.restore(archive, destination)
    assert not destination.exists()


def test_empty_workbook_has_no_fabricated_readable_body(tmp_path):
    path = tmp_path / 'empty.xlsx'
    workbook = openpyxl.Workbook()
    workbook.save(path)
    assert preprocessor.extract_text_from_xlsx(str(path)) == ''


def test_docx_cell_span_rejected_before_document_decoder(tmp_path, monkeypatch):
    import docx
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'small'
    original = io.BytesIO()
    document.save(original)
    output = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == 'word/document.xml':
                data = data.replace(b'<w:gridSpan w:val="2"', b'<w:gridSpan w:val="999999999"')
            target.writestr(entry, data)
    path = tmp_path / 'span.docx'
    path.write_bytes(output.getvalue())
    monkeypatch.setattr(docx, 'Document', lambda *_a, **_k: pytest.fail('decoder reached for enormous span'))
    with pytest.raises(input_safety.UnsafeDocument):
        preprocessor.extract_text_from_docx(str(path))


def test_hwpx_entities_are_not_returned_as_fallback_text(tmp_path):
    path = tmp_path / 'entities.hwpx'
    path.write_bytes(package({'Contents/section0.xml': '<!DOCTYPE p [<!ENTITY x "small">]><p>&x;</p>'}))
    with pytest.raises(input_safety.UnsafeDocument):
        preprocessor.extract_text_from_hwpx(str(path))


def test_multiline_filename_cannot_inject_readable_body(tmp_path):
    path = tmp_path / 'blank.txt'
    path.write_text('   ')
    text, kind = preprocessor.extract_text_from_file(str(path), 'blank\n상호: 가짜회사\n합계금액: 999999\n.txt')
    assert kind == 'txt'
    assert '\n' not in text
    assert text.partition('\n')[2] == ''
