import copy
import json

import pytest

from backend.evaluation_confirmation import apply_human_decisions
from backend.prediction_artifacts import write_json, file_hash
from backend.scripts.import_reconciled_labels import import_review, read_wording_resolutions
from backend.scripts.prepare_test_review import page_builder, workbook_from_comparison
from backend.evaluation_review import compare_labels
from test_evaluation import confirmation_example, source_example, independent_example


def resolution(reason):
    return {"review_id": "1/7/total_amount", "kind": "QUESTION_WORDING_ERROR_LABEL_UNCHANGED",
        "adjudicator": "unit-fixture", "human_observation": reason,
        "source_sha256": "a" * 64, "original_question": "unit question erroneously says no label",
        "corrected_question": "unit question: number is covered", "resolved_answer": {"state": "REVIEW", "value": None},
        "resolution_reason": "unit observation supports the unchanged REVIEW label"}


def test_wording_error_keeps_original_dissent_and_does_not_invent_a_value():
    manifest, reference, comparison, decisions = confirmation_example()
    decision = decisions["1/7/total_amount"]
    decision.update(confirmation="수정 필요", evidence="unit fixture: label exists but value is covered")
    with pytest.raises(ValueError):
        apply_human_decisions(manifest, reference, comparison, decisions, "unit-human")
    decision["wording_resolution"] = resolution(decision["evidence"])
    labels = apply_human_decisions(manifest, reference, comparison, decisions, "unit-human")
    answer = labels["documents"]["7"]["fields"]["total_amount"]
    assert answer["state"] == "REVIEW" and answer["value"] is None
    assert answer["review"]["original_human_decision"] == "수정 필요"
    assert decision["confirmation"] == "수정 필요" and reference["review"]["confirmed"] is False


@pytest.mark.parametrize("change", ["numeric_value", "different_reason", "not_review"])
def test_wording_exception_cannot_bypass_real_label_changes(change):
    manifest, reference, comparison, decisions = confirmation_example()
    decision = decisions["1/7/total_amount"]
    decision.update(confirmation="수정 필요", evidence="unit fixture covered value")
    decision["wording_resolution"] = resolution(decision["evidence"])
    if change == "numeric_value":
        decision["value"] = 110
    elif change == "different_reason":
        decision["wording_resolution"]["human_observation"] = "different observation"
    else:
        decision["state"] = "ABSENT"
    with pytest.raises(ValueError):
        apply_human_decisions(manifest, reference, comparison, decisions, "unit-human")


def test_review_question_distinguishes_missing_item_from_unreadable_value():
    builder = page_builder()
    question, hint = builder.field_q({"I": "REVIEW", "J": None}, {"doctype": {"7": "견적서"}}, "7", "total_amount")
    assert "이름의 금액이 없다" not in question
    assert "숫자를 읽을 수 없거나" in question and "가려져" in question and "도장" in hint


@pytest.mark.parametrize("rule", ["ORIGINAL_FILE", "DOCUMENT_PROCESSING", "DOCUMENT_TYPE", "AMOUNT_COMPONENTS", "BUSINESS_NO_MATCH"])
def test_not_applicable_question_does_not_substitute_a_positive_source_fact(rule):
    context = {"doctype": {}, "fields": {"7": {field: {"state": "VALUE", "value": value}
        for field, value in (("supply_amount", 100), ("tax_amount", 10), ("total_amount", 110))}}}
    question, hint = page_builder().check_q({"I": "N/A", "J": None}, context, rule, "7")
    assert "해당 없음(N/A)" in question and "PASS와 구분" in question
    assert "원문에서 직접 확인한 사실" in hint
    assert "잘 모르겠어요" in question


def test_review_question_does_not_tell_human_to_accept_the_draft():
    builder = page_builder()
    item = {"A": "1/7/total_amount", "C": "total_amount", "H": "해석 확인", "I": "REVIEW", "J": None,
        "M": "unit fixture: unreadable value", "sheet": "필드확인", "row": 5}
    cases = {"1": {"doctype": {"7": "견적서"}}}
    question = builder.make_items([item], cases)[0]['q']
    assert "기준대로면" not in question and "설명대로면" not in question
    assert "애매하면 잘 모르겠어요" in question


def test_import_keeps_raw_human_json_and_excel_dissent_with_bound_resolution(tmp_path):
    import openpyxl
    manifest, reference = source_example()
    manifest["contracts"][0]["split"] = "test"
    reference["split"] = "test"
    source = tmp_path / "fiction.txt"
    source.write_text("unit fixture: label exists, numeric value covered")
    manifest["contracts"][0]["documents"][0].update(path=source.name, sha256=file_hash(source))
    write_json(tmp_path / "manifest.json", manifest)
    write_json(tmp_path / "reference_labels_test.json", reference)
    write_json(tmp_path / "reference_labels_test.freeze.json", {"manifest_sha256": file_hash(tmp_path / "manifest.json"),
        "labels_sha256": file_hash(tmp_path / "reference_labels_test.json")})
    write_json(tmp_path / "ground_truth_test.json", reference)
    replies = tmp_path / "replies"
    reply = independent_example(manifest, reference)
    write_json(replies / "fixture.json", reply)
    comparison = compare_labels(manifest, reference, reply)
    comparison["hashes"] = {"manifest": file_hash(tmp_path / "manifest.json"),
        "reference": file_hash(tmp_path / "reference_labels_test.json"),
        "independent_replies": {"fixture.json": file_hash(replies / "fixture.json")}}
    write_json(tmp_path / "comparison.json", comparison)
    workbook = tmp_path / "review.xlsx"
    workbook_from_comparison(tmp_path, tmp_path / "comparison.json", workbook)
    initial_hash = file_hash(workbook)
    book = openpyxl.load_workbook(workbook)
    book["안내"]["B12"] = "unit-human"
    items = []
    for name in ("필드확인", "검수확인"):
        for row in range(5, book[name].max_row + 1):
            rid = book[name].cell(row, 1).value
            if not rid:
                continue
            decision = "수정 필요" if rid == "1/7/total_amount" else "확인"
            reason = "unit fixture: label exists but value is covered" if decision == "수정 필요" else "unit direct observation"
            book[name].cell(row, 11, decision)
            book[name].cell(row, 12, reason)
            items.append({"id": rid, "decision": decision, "reason": reason, "reviewer": "unit-human",
                "reviewed_at": "2026-01-01T12:00:00+09:00", "sheet": name, "row": row})
    book.save(workbook)
    book.close()
    raw = tmp_path / "raw-human.json"
    write_json(raw, {"reviewer": "unit-human", "items": items, "source_xlsx_sha256": initial_hash})
    raw_hash = file_hash(raw)
    wording = resolution(next(i["reason"] for i in items if i["decision"] == "수정 필요"))
    wording["source_sha256"] = file_hash(source)
    resolved = tmp_path / "wording.json"
    record = {"format": "human-question-wording-resolution-v1", "human_result_sha256": raw_hash,
        "source_xlsx_sha256": initial_hash, "items": [wording]}
    write_json(resolved, record)
    bad = copy.deepcopy(record)
    bad["items"][0]["source_sha256"] = "different-source"
    write_json(tmp_path / "bad.json", bad)
    with pytest.raises(ValueError):
        read_wording_resolutions(tmp_path / "bad.json", raw, manifest)
    labels = import_review(workbook, tmp_path, tmp_path / "comparison.json", replies, "unit-human", raw,
        split="test", resolution_path=resolved)
    assert labels["review"]["prompt_wording_error_count"] == 1
    assert labels["review"]["original_human_decision_counts"]["수정 필요"] == 1
    assert file_hash(raw) == raw_hash
    answer = labels["documents"]["7"]["fields"]["total_amount"]
    assert answer["review"]["original_human_decision"] == "수정 필요" and answer["value"] is None
