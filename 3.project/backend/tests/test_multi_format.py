import os
import tempfile
import zipfile
import pytest
import docx
import openpyxl
from backend.app.ai.preprocessor import extract_text_from_file
from backend.app.ai import analyze_document
from backend.app.schemas.common import DocumentType


def test_docx_parsing_and_classification():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = os.path.join(tmp_dir, "01_표준_외주용역계약서.docx")
        doc = docx.Document()
        doc.add_heading("외주 용역 표준 계약서", 0)
        doc.add_paragraph("본 계약은 발주사 주식회사 갑과 수주사 주식회사 테스트소프트 간에 체결한다.")
        doc.add_paragraph("제 1 조 (목적): 본 계약은 웹 애플리케이션 개발 용역의 수행에 관한 사항을 정한다.")
        doc.add_paragraph("계약금액: 15,000,000 원 (VAT 별도)")
        doc.add_paragraph("계약일자: 2026-03-15")
        
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "상호명"
        table.rows[0].cells[1].text = "(주)테스트소프트"
        table.rows[1].cells[0].text = "사업자등록번호"
        table.rows[1].cells[1].text = "123-45-67890"
        
        doc.save(file_path)

        text, ext = extract_text_from_file(file_path, "01_표준_외주용역계약서.docx")
        assert ext == "docx"
        assert "외주 용역 표준 계약서" in text
        assert "15,000,000" in text

        analysis, raw_t, masked_t = analyze_document(file_path, file_name="01_표준_외주용역계약서.docx")
        assert analysis.document_type == DocumentType.CONTRACT
        assert analysis.confidence >= 0.80
        assert analysis.fields.amount == 15000000


def test_xlsx_parsing_and_classification():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = os.path.join(tmp_dir, "02_개발용역_견적서.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "견적서"
        ws.append(["견 적 서"])
        ws.append(["귀하 아래와 같이 견적합니다."])
        ws.append(["견적일자", "2026-03-20"])
        ws.append(["공급자", "(주)소프트웨어"])
        ws.append(["사업자등록번호", "234-56-78901"])
        ws.append(["총 견적금액", "8,000,000원"])
        ws.append(["품목", "수량", "단가", "견적금액"])
        ws.append(["백엔드 개발", "1", "8,000,000원", "8,000,000원"])
        wb.save(file_path)

        text, ext = extract_text_from_file(file_path, "02_개발용역_견적서.xlsx")
        assert ext == "xlsx"
        assert "견 적 서" in text
        assert "8,000,000" in text

        analysis, raw_t, masked_t = analyze_document(file_path, file_name="02_개발용역_견적서.xlsx")
        assert analysis.document_type == DocumentType.ESTIMATE
        assert analysis.confidence >= 0.80
        assert analysis.fields.amount == 8000000


def test_hwpx_parsing_and_classification():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = os.path.join(tmp_dir, "06_최종_검수확인서.hwpx")
        xml_content = """<?xml version="1.0" encoding="UTF-8"?>
        <hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
          <hp:p><hp:t>용역 완료 및 검수확인서</hp:t></hp:p>
          <hp:p><hp:t>계약상대자: (주)테스트소프트</hp:t></hp:p>
          <hp:p><hp:t>상기 용역이 과업지시서에 따라 적합하게 완료되었음을 확인하고 이에 검수합니다.</hp:t></hp:p>
          <hp:p><hp:t>검수일자: 2026-04-30</hp:t></hp:p>
          <hp:p><hp:t>검수자: 홍길동 팀장</hp:t></hp:p>
          <hp:p><hp:t>검수금액: 15,000,000 원</hp:t></hp:p>
        </hs:sec>"""
        with zipfile.ZipFile(file_path, "w") as zf:
            zf.writestr("Contents/section0.xml", xml_content.encode("utf-8"))

        text, ext = extract_text_from_file(file_path, "06_최종_검수확인서.hwpx")
        assert ext == "hwpx"
        assert "검수확인서" in text
        assert "완료되었음을 확인" in text

        analysis, raw_t, masked_t = analyze_document(file_path, file_name="06_최종_검수확인서.hwpx")
        assert analysis.document_type == DocumentType.INSPECTION_CONFIRMATION
        assert analysis.confidence >= 0.80


def test_csv_parsing():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = os.path.join(tmp_dir, "05_세금계산서내역.csv")
        content = "전자세금계산서\n공급자,(주)테스트\n공급받는자,(주)발주사\n합계금액,5500000\n작성일자,2026-04-10"
        with open(file_path, "w", encoding="utf-8-sig") as f:
            f.write(content)

        text, ext = extract_text_from_file(file_path, "05_세금계산서내역.csv")
        assert ext == "csv"
        assert "전자세금계산서" in text

        analysis, raw_t, masked_t = analyze_document(file_path, file_name="05_세금계산서내역.csv")
        assert analysis.document_type == DocumentType.TAX_INVOICE
