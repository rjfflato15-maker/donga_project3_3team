"""Safety cases on hand-authored sources; no model calls or evaluation answers."""
import json
from types import SimpleNamespace

import pytest
import requests

from backend.app.ai.grounding import guard_proposal, supplementation_targets, source_pages
from backend.app.ai.llm_supplement import apply_proposals, supplement_document
from backend.app.ai.local_llm import LLMSettings, OllamaClient, valid_answer


def page(text, number=1):
    return [{"page": number, "location_kind": "page" if number else "document_block", "text": text}]


def proposal(value, quote, line=1, number=1, state="VALUE"):
    return {"state": state, "value": value, "page": number, "line": line, "quote": quote}


def test_guard_rejects_recipient_number_even_when_it_exists_in_source():
    text = "공급자\n사업자등록번호: 000-00-01014\n공급받는자\n사업자등록번호: 000-00-11017"
    result = guard_proposal("business_registration_no", proposal("000-00-11017", text.splitlines()[3], 4), page(text), "tax_invoice")
    assert not result["accepted"] and "WRONG_PARTY" in result["reasons"]


def test_guard_rejects_value_from_different_field_label():
    text = "공급가액: 100원\n세액: 10원\n합계금액: 110원"
    result = guard_proposal("tax_amount", proposal("100", "공급가액: 100원"), page(text), "tax_invoice")
    assert "FIELD_LABEL_OR_VALUE_MISMATCH" in result["reasons"]


@pytest.mark.parametrize("text,value", [("합계(VAT 제외): 100원", "100"), ("합계금액: 2026-09-01", "2026"), ("세액: 10%", "10")])
def test_guard_does_not_treat_vat_exclusion_dates_or_rates_as_total_amount(text, value):
    field = "tax_amount" if text.startswith("세액") else "total_amount"
    result = guard_proposal(field, proposal(value, text), page(text), "estimate")
    assert not result["accepted"]


def test_guard_rejects_invented_value_and_label_only_quote():
    text = "합계금액: 110원"
    invented = guard_proposal("total_amount", proposal("111", text), page(text), "estimate")
    assert not invented["accepted"] and "VALUE_NOT_IN_QUOTE" in invented["reasons"]
    label_only = guard_proposal("total_amount", proposal("110", "합계금액"), page(text), "estimate")
    assert "VALUE_NOT_IN_QUOTE" in label_only["reasons"]
    missing_company_quote = guard_proposal("company_name", proposal("테스트 회사", "UNKNOWN"), page("상호: 테스트 회사"), "business_registration")
    assert not missing_company_quote["accepted"] and "VALUE_NOT_IN_QUOTE" in missing_company_quote["reasons"]


def test_guard_does_not_guess_unlabeled_reference_amounts():
    text = "견적서\n참고값 A: 100원\n참고값 B: 110원"
    result = guard_proposal("total_amount", proposal("110", text.splitlines()[2], 3), page(text), "estimate")
    assert not result["accepted"]
    abstention = guard_proposal("total_amount", proposal(None, "", None, None, "REVIEW"), page(text), "estimate")
    assert abstention["accepted"]
    scope = guard_proposal("amount_scope", proposal(None, "", None, None, "REVIEW"), page(text), "estimate")
    assert scope["accepted"]


def test_ocr_dot_separator_is_grounded_without_rewriting_values():
    text = "공급자\n상호. 테스트 회사\n사업자등록번호. 000-00-01014\n공급가액. 100원\n세액. 10원\n합계금액. 110원\n발행일. 2026-09-01"
    for field, value, line in (("company_name", "테스트 회사", 2), ("business_registration_no", "000-00-01014", 3),
                               ("total_amount", "110", 6), ("issue_date", "2026-09-01", 7)):
        result = guard_proposal(field, proposal(value, text.splitlines()[line - 1], line), page(text), "tax_invoice")
        assert result["accepted"], result
        assert result["anchor"]["quote"] == text.splitlines()[line - 1]
    result = guard_proposal("total_amount", proposal("5", "합계금액.5원"), page("합계금액.5원"), "estimate")
    assert not result["accepted"]


def test_guard_rejects_multiple_labeled_candidates_and_unverified_page():
    text = "합계금액: 110원\n합계금액: 120원"
    result = guard_proposal("total_amount", proposal("110", text.splitlines()[0]), page(text), "estimate")
    assert "AMBIGUOUS" in result["reasons"]
    result = guard_proposal("total_amount", proposal("110", text.splitlines()[0], number=None), page(text, None), "estimate")
    assert "PAGE_UNVERIFIED" in result["reasons"]


def test_printed_arithmetic_error_is_preserved_for_rule_validation():
    text = "공급가액: 100원\n세액: 10원\n합계금액: 111원"
    result = guard_proposal("total_amount", proposal("111", text.splitlines()[2], 3), page(text), "tax_invoice")
    assert result["accepted"] and "SOURCE_ARITHMETIC_MISMATCH" in result["warnings"]
    assert result["anchor"]["page"] == 1 and result["anchor"]["line"] == 3


def test_failed_llm_retains_rule_values_and_manual_review_reason():
    fields = {"total_amount": 123, "amount_scope": "FULL"}
    result, review = apply_proposals(fields, {"ok": False, "error": "TIMEOUT", "model": "qwen3:latest"}, [], "estimate", ["total_amount"])
    assert result == fields and review["error"] == "TIMEOUT" and review["manual_review_required"]


def test_successful_rule_fields_skip_llm_entirely():
    text = "사업자등록증\n상호: 테스트 회사\n사업자등록번호: 000-00-01014\n발급일: 2026-09-01"
    fields = {"company_name": "테스트 회사", "business_registration_no": "000-00-01014", "issue_date": "2026-09-01"}
    client = SimpleNamespace(extract=lambda *a: pytest.fail("unnecessary model call"))
    result, review = supplement_document("virtual.txt", text, fields, "business_registration", client)
    assert result == fields and review["status"] == "NOT_NEEDED"


def test_targets_detect_wrong_party_without_requesting_successful_amount_again():
    text = "공급자\n사업자등록번호: 000-00-01014\n공급받는자\n사업자등록번호: 000-00-11017\n합계금액: 110원"
    fields = {"business_registration_no": "000-00-11017", "total_amount": 110, "amount_scope": "FULL"}
    targets = supplementation_targets(fields, page(text), "tax_invoice")
    assert "business_registration_no" in targets and "total_amount" not in targets


@pytest.mark.parametrize("settings", [LLMSettings(url="https://ollama.com"), LLMSettings(url="http://remote.invalid:11434"),
                                     LLMSettings(model="qwen3:cloud"), LLMSettings(model="gemma4:latest")])
def test_client_refuses_remote_or_unapproved_models(settings):
    with pytest.raises(ValueError):
        OllamaClient(settings)


def test_json_settings_and_timeout_retry_do_not_log_source(tmp_path, monkeypatch, capsys):
    from backend.app.ai import local_llm
    (tmp_path / ".ollama").mkdir()
    (tmp_path / ".ollama/server.json").write_text('{"disable_ollama_cloud":true}', encoding="utf-8")
    monkeypatch.setattr(local_llm.Path, "home", lambda: tmp_path)

    class FakeSession:
        trust_env = True
        count = 0
        def get(self, *args, **kwargs):
            return SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"models": [{"name": "qwen3:latest", "size": 100, "digest": "local"}]})
        def post(self, *args, **kwargs):
            self.count += 1
            assert kwargs["allow_redirects"] is False
            assert kwargs["json"]["options"]["temperature"] == 0
            assert kwargs["json"]["think"] is False
            assert kwargs["json"]["format"]["type"] == "object"
            raise requests.Timeout("PRIVATE_SOURCE_SHOULD_NOT_BE_LOGGED")
    session = FakeSession()
    result = OllamaClient(session=session).extract(page("PRIVATE_DOCUMENT"), ["company_name"], "contract")
    assert result["error"] == "TIMEOUT" and result["attempts"] == 2 and session.trust_env is False
    assert capsys.readouterr().out == ""


def test_cloud_config_is_required_before_sending_document(tmp_path, monkeypatch):
    from backend.app.ai import local_llm
    monkeypatch.setattr(local_llm.Path, "home", lambda: tmp_path)
    session = SimpleNamespace(get=lambda *a, **kw: pytest.fail("network before cloud config check"),
                              post=lambda *a, **kw: pytest.fail("source sent before cloud config check"))
    run = OllamaClient(session=session).extract(page("PRIVATE_DOCUMENT"), ["company_name"], "contract")
    assert not run["ok"] and run["attempts"] == 0 and run["error"] == "CLOUD_DISABLE_CONFIG_REQUIRED"


def test_ocr_fallback_does_not_print_private_exception_contents(monkeypatch, capsys):
    import sys
    from PIL import Image
    from backend.app.ai import preprocessor
    def failed_ocr(*args, **kwargs):
        raise RuntimeError("PRIVATE_DOCUMENT_PATH_OR_TEXT")
    monkeypatch.setattr(preprocessor, "_safe_winocr_recognize", failed_ocr)
    monkeypatch.setitem(sys.modules, "pytesseract", SimpleNamespace(image_to_string=lambda *a, **kw: ""))
    with Image.new("RGB", (10, 10)) as source:
        assert preprocessor.run_ocr_on_image(source) == ""
    assert "PRIVATE_DOCUMENT" not in capsys.readouterr().out


def test_guarded_and_unguarded_apply_exact_same_response_with_different_outcomes():
    text = "참고값 A: 100원\n참고값 B: 110원"
    run = {"ok": True, "proposals": {"total_amount": proposal("110", text.splitlines()[1], 2)}}
    raw, _ = apply_proposals({}, run, page(text), "estimate", ["total_amount"], guarded=False)
    guarded, review = apply_proposals({}, run, page(text), "estimate", ["total_amount"])
    assert raw["total_amount"] == 110 and "total_amount" not in guarded
    assert review["decisions"][0]["human_confirmed"] is False


def test_docx_paragraph_and_table_evidence_preserves_unverified_page_layout(tmp_path):
    from docx import Document
    from backend.app.ai.preprocessor import extract_text_from_docx
    document = Document()
    document.add_paragraph("공급자")
    document.add_paragraph("상호: 테스트 회사")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "합계금액"
    table.cell(0, 1).text = "110원"
    path = tmp_path / "source.docx"
    document.save(path)
    text = extract_text_from_docx(path)
    pages = source_pages(path, text)
    company = guard_proposal("company_name", proposal("테스트 회사", text.splitlines()[1], 2, None), pages, "tax_invoice")
    amount = guard_proposal("total_amount", proposal("110", text.splitlines()[2], 3, None), pages, "tax_invoice")
    assert company["accepted"] and amount["accepted"]
    assert company["anchor"]["positions"][0]["paragraph"] == 2
    assert amount["anchor"]["positions"][-1]["table"] == 1
    assert amount["anchor"]["positions"][-1]["cell"] == 2
    for decision in (company, amount):
        assert decision["anchor"]["page"] is None
        assert decision["anchor"]["text_verified"] is True
        assert decision["anchor"]["page_layout_verified"] is False
    tampered = source_pages(path, text.replace("110원", "111원"))
    rejected = guard_proposal("total_amount", proposal("111", "합계금액\t111원", 3, None), tampered, "tax_invoice")
    assert not rejected["accepted"] and "DOCX_LOCATION_UNVERIFIED" in rejected["reasons"]


def test_docx_table_does_not_turn_recipient_into_supplier(tmp_path):
    from docx import Document
    from backend.app.ai.preprocessor import extract_text_from_docx
    document = Document()
    document.add_paragraph("공급자")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "공급받는자"
    table.cell(1, 0).text = "상호"
    table.cell(1, 1).text = "수신 회사"
    path = tmp_path / "recipient.docx"
    document.save(path)
    text = extract_text_from_docx(path)
    result = guard_proposal("company_name", proposal("수신 회사", "상호\t수신 회사", 3, None), source_pages(path, text), "tax_invoice")
    assert not result["accepted"] and "WRONG_PARTY" in result["reasons"]


def test_docx_multiline_merged_cells_retain_native_positions(tmp_path):
    from docx import Document
    from backend.app.ai.preprocessor import extract_text_from_docx
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "합계금액\n110원"
    path = tmp_path / "merged.docx"
    document.save(path)
    text = extract_text_from_docx(path)
    result = guard_proposal("total_amount", proposal("110", "110원", 2, None), source_pages(path, text), "estimate")
    assert result["accepted"] and result["anchor"]["label_line"] == 1
    assert result["anchor"]["positions"][0]["row"] == 1


@pytest.mark.parametrize("reported_line", [1, 3])
def test_one_line_offset_requires_same_exact_quote_label_and_value(reported_line):
    text = "비고\n합계금액: 110원\n작성일: 2026-09-01"
    result = guard_proposal("total_amount", proposal("110", "합계금액: 110원", reported_line), page(text), "estimate")
    assert result["accepted"] and result["anchor"]["line"] == 2
    assert result["proposal"]["line"] == reported_line
    assert "LINE_OFFSET_RESOLVED" in result["warnings"]


def test_offset_does_not_cross_party_wrong_page_or_two_lines():
    text = "공급자\n사업자등록번호: 000-00-01014\n공급받는자"
    result = guard_proposal("business_registration_no", proposal("000-00-01014", text.splitlines()[1], 3), page(text), "tax_invoice")
    assert not result["accepted"]
    result = guard_proposal("total_amount", proposal("110", "합계금액: 110원", number=2), page("합계금액: 110원"), "estimate")
    assert not result["accepted"] and "PAGE_UNVERIFIED" in result["reasons"]
    result = guard_proposal("total_amount", proposal("110", "합계금액: 110원", 1), page("비고\n비고\n합계금액: 110원"), "estimate")
    assert not result["accepted"] and "EVIDENCE_NOT_FOUND" in result["reasons"]
    result = guard_proposal("tax_amount", proposal("110", "합계금액: 110원", 1), page("비고\n합계금액: 110원"), "estimate")
    assert not result["accepted"]


@pytest.mark.parametrize("text,value,accepted", [("합계금액. 1 700,000원", "1700000", True),
    ("합계금액: 1 700 000원", "1700000", True), ("합계금액: 100 200", "100200", False),
    ("합계금액: 1 70,000원", "170000", False), ("합계금액: 100,20원", "10020", False)])
def test_space_grouped_currency_is_narrowly_normalized(text, value, accepted):
    result = guard_proposal("total_amount", proposal(value, text), page(text), "estimate")
    assert result["accepted"] is accepted


def test_ocr_bullet_is_only_stripped_after_a_recognized_label():
    result = guard_proposal("company_name", proposal("테스트•회사", "예금주• 테스트•회사"), page("예금주• 테스트•회사"), "bank_account")
    assert result["accepted"]
    text = "발주자: 수신 회사\n=(크): 테스트 회사"
    result = guard_proposal("company_name", proposal("테스트 회사", text.splitlines()[1], 2), page(text), "contract")
    assert not result["accepted"]


def test_missing_scope_anchor_is_derived_only_from_unique_labeled_evidence():
    text = "견적서\n합계금액: 110원"
    result = guard_proposal("amount_scope", proposal("FULL", "", None, None), page(text), "estimate")
    assert result["accepted"] and result["anchor"]["quote"] == "견적서"
    assert "CODE_ANCHORED_SCOPE" in result["warnings"]
    for text in ("견적서\n참고값 A: 100원\n참고값 B: 110원",):
        result = guard_proposal("amount_scope", proposal("FULL", "", None, None), page(text), "estimate")
        assert not result["accepted"]
    # A payment stage word does not change the aggregate estimate's scope.
    text = "견적서\n합계금액: 110원\n잔금"
    assert guard_proposal("amount_scope", proposal("FULL", "", None, None), page(text), "estimate")["accepted"]


def test_original_incorrect_scope_patterns_and_absent_reference_values_stay_rejected():
    text = "검수확인서\n합계금액: 110원"
    result = guard_proposal("amount_scope", proposal(None, "", None, None, "ABSENT"), page(text), "inspection_confirmation")
    assert not result["accepted"] and "UNSUPPORTED_ABSTENTION" in result["reasons"]
    text = "견적서\n참고값 A: 100원\n참고값 B: 110원\n금액의 적용 범위가 표시되지 않음"
    result = guard_proposal("amount_scope", proposal("PARTIAL", text.splitlines()[3], 4), page(text), "estimate")
    assert not result["accepted"]
    run = {"ok": True, "proposals": {"total_amount": proposal(None, "", None, None, "ABSENT")}}
    fields, review = apply_proposals({}, run, page(text), "estimate", ["total_amount"])
    assert "total_amount" not in fields and review["guard_review_required"]
    assert "AMBIGUOUS_UNLABELED_AMOUNT" in review["decisions"][0]["reasons"]


def test_machine_grounding_does_not_claim_final_human_confirmation():
    text = "合計\n합계금액: 110원"
    run = {"ok": True, "proposals": {"total_amount": proposal("110", text.splitlines()[1], 2)}}
    _, review = apply_proposals({}, run, page(text), "estimate", ["total_amount"])
    assert not review["guard_review_required"]
    assert review["manual_review_required"] and not review["human_confirmed"]


def test_schema_rejects_malformed_states_and_anchor_types():
    assert not valid_answer(proposal("100", "", state="REVIEW"))
    assert not valid_answer(proposal("100", None))
    assert not valid_answer(proposal("100", "합계: 100", line=True))
    assert valid_answer(proposal(None, "", None, None, "REVIEW"))


def test_ai_provenance_survives_storage_and_human_correction(operator_client, monkeypatch):
    from backend.app.ai import llm_supplement
    monkeypatch.setenv("APP_LLM_ENABLED", "true")
    review = {"status": "FAILED", "error": "TIMEOUT", "manual_review_required": True, "human_confirmed": False,
              "decisions": [{"field": "company_name", "applied": False, "reasons": ["WRONG_PARTY"]}]}
    monkeypatch.setattr(llm_supplement, "supplement_document", lambda path, raw, fields, dtype: (fields, review))
    contract = operator_client.post("/api/contracts", json={"title": "검수 테스트", "vendor_name": "테스트회사",
        "business_number": "000-00-01014", "contract_amount": 110, "amount_basis": "TOTAL", "payment_mode": "FULL"}).json()
    cid = contract["contract_id"]
    response = operator_client.post(f"/api/contracts/{cid}/documents", files=[("files", ("문서.txt", "계약서\n합계금액: 110원".encode(), "text/plain"))])
    assert response.status_code == 200
    did = response.json()[0]["document_id"]
    doc = operator_client.get(f"/api/documents/{did}").json()
    assert doc["analysis"]["ai_review"] == review
    revision = operator_client.get(f"/api/contracts/{cid}").json()["workflow"]["revision"]
    corrected = operator_client.put(f"/api/documents/{did}/fields", json={"fields": doc["analysis"]["fields"],
        "document_type": doc["document_type"], "revision": revision, "note": "원문 직접 확인",
        "confirmed_fields": ["company_name"]})
    assert corrected.status_code == 200, corrected.text
    saved = corrected.json()["analysis"]["ai_review"]
    assert saved["human_confirmed"] and not saved["manual_review_required"]
    assert saved["decisions"] == [{**decision, "human_confirmed": True} for decision in review["decisions"]]
    assert saved["status"] == "FAILED" and saved["error"] == "TIMEOUT"
