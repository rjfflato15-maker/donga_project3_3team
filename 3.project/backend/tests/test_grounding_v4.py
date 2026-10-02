"""Independent safety examples; no v1 test or test2 fixture is imported."""
import pytest
from backend.app.ai import grounding as v3
from backend.app.ai.grounding_v4 import guard_proposal, GUARD_VERSION
from backend.app.ai.llm_supplement import apply_proposals, supplement_document


def page(text, number=1):
    return [{'page':number,'location_kind':'page','text':text}]


def proposal(value, quote, line=1, number=1, state='VALUE'):
    return {'state':state,'value':value,'quote':quote,'line':line,'page':number}


@pytest.mark.parametrize('field,text,value',[
    ('supply_amount','세전 대금: 1,000원','1000'),
    ('tax_amount','부가세액: 100원','100'),
    ('total_amount','총 청구액: 1,100원','1100'),
])
def test_explicit_dev2_labels(field,text,value):
    answer=proposal(value,text)
    assert guard_proposal(field,answer,page(text),'estimate')['accepted']
    assert not v3.guard_proposal(field,answer,page(text),'estimate')['accepted']


@pytest.mark.parametrize('quote,line',[('공급가액',1),('1,000원',2)])
def test_adjacent_table_label_and_value_have_separate_verified_anchors(quote,line):
    d=guard_proposal('supply_amount',proposal('1000',quote,line),page('공급가액\n1,000원\n세액\n100원'),'estimate')
    assert d['accepted']
    assert d['anchor']['line']==2
    assert d['anchor']['label_line']==1
    assert d['anchor']['label_anchor']['quote']=='공급가액'
    assert d['anchor']['quote']=='1,000원'
    assert not d['human_confirmed']


def test_multiple_fields_in_one_line_do_not_extend_company_name():
    text='공급자 상호: 한결소프트\t사업자등록번호: 123-45-67891'
    d=guard_proposal('company_name',proposal('한결소프트',text),page(text),'tax_invoice')
    assert d['accepted']
    assert d['anchor']['party']=='supplier'


def test_buyer_number_still_rejected():
    text='공급자\n사업자등록번호: 123-45-67891\n공급받는자\n사업자등록번호: 111-11-11119'
    d=guard_proposal('business_registration_no',proposal('1111111119',text.splitlines()[3],4),page(text),'tax_invoice')
    assert not d['accepted'] and 'WRONG_PARTY' in d['reasons']


def test_offset_cannot_cross_buyer_boundary():
    text='공급자\n사업자등록번호: 123-45-67891\n공급받는자'
    d=guard_proposal('business_registration_no',proposal('1234567891',text.splitlines()[1],3),page(text),'tax_invoice')
    assert not d['accepted'] and 'EVIDENCE_NOT_FOUND' in d['reasons']


def test_two_parties_in_one_line_are_distinguished_at_label():
    text='공급자 상호: 한결소프트\t공급받는자 상호: 북쪽회사'
    assert guard_proposal('company_name',proposal('한결소프트',text),page(text),'tax_invoice')['accepted']
    d=guard_proposal('company_name',proposal('북쪽회사',text),page(text),'tax_invoice')
    assert not d['accepted'] and 'WRONG_PARTY' in d['reasons']


def test_made_up_number_is_not_repaired_from_arithmetic():
    text='공급가액: 1,000원\n세액: 100원'
    d=guard_proposal('total_amount',proposal('1100',text.splitlines()[0]),page(text),'estimate')
    assert not d['accepted'] and 'FIELD_LABEL_OR_VALUE_MISMATCH' in d['reasons']


def test_same_number_under_wrong_field_does_not_pass():
    text='세액: 1,000원'
    assert not guard_proposal('total_amount',proposal('1000',text),page(text),'estimate')['accepted']


def test_reference_values_stay_review():
    text='견적서\n참고값 A: 1,000원\n참고값 B: 1,100원'
    assert not guard_proposal('total_amount',proposal('1100',text.splitlines()[2],3),page(text),'estimate')['accepted']
    review=proposal(None,'',None,None,'REVIEW')
    assert guard_proposal('total_amount',review,page(text),'estimate')['accepted']
    absent=proposal(None,'',None,None,'ABSENT')
    assert 'AMBIGUOUS_UNLABELED_AMOUNT' in guard_proposal('total_amount',absent,page(text),'estimate')['reasons']


def test_distinct_labelled_totals_are_ambiguous():
    text='금액 안내\n합계금액: 1,100원\n총 청구액: 2,200원'
    d=guard_proposal('total_amount',proposal('1100',text.splitlines()[1],2),page(text),'estimate')
    assert not d['accepted'] and 'AMBIGUOUS' in d['reasons']


@pytest.mark.parametrize('answer',[
    proposal('1100','합계금액: 1,100원',number=2),
    proposal('1100','합계금액: 1,200원'),
    proposal('1100','합계금액: 1,100원',line=4),
])
def test_wrong_page_quote_or_large_offset_fails(answer):
    assert not guard_proposal('total_amount',answer,page('합계금액: 1,100원'),'estimate')['accepted']


@pytest.mark.parametrize('text,value,field',[
    ('사업자등록번호: 123-45-67890','1234567890','business_registration_no'),
    ('총 청구액: -1,100원','-1100','total_amount'),
    ('총 청구액: 10%','10','total_amount'),
    ('발행일: 2025-02-30','2025-02-30','issue_date'),
])
def test_format_and_checksum_remain_enforced(text,value,field):
    assert not guard_proposal(field,proposal(value,text),page(text),'business_registration')['accepted']


def test_printed_arithmetic_error_is_preserved_and_flagged():
    text='공급가액: 1,000원\n세액: 100원\n총 청구액: 1,200원'
    d=guard_proposal('total_amount',proposal('1200',text.splitlines()[2],3),page(text),'estimate')
    assert d['accepted'] and 'SOURCE_ARITHMETIC_MISMATCH' in d['warnings']
    assert d['proposal']['value']=='1200'


def test_absence_does_not_erase_unverified_financial_field():
    run={'ok':True,'proposals':{'total_amount':proposal(None,'',None,None,'ABSENT')}}
    fields,review=apply_proposals({'total_amount':1000},run,page('견적서\n확인 인'),'estimate',['total_amount'],guard_policy='v4')
    assert fields['total_amount']==1000
    assert review['guard_review_required'] and review['guard_version']==GUARD_VERSION
    assert 'UNVERIFIED_ABSENCE' in review['decisions'][0]['reasons']


def test_v3_default_and_v4_option_are_separate():
    text='총 청구액: 1,100원';run={'ok':True,'proposals':{'total_amount':proposal('1100',text)}}
    a,old=apply_proposals({'total_amount':None},run,page(text),'estimate',['total_amount'])
    b,new=apply_proposals({'total_amount':None},run,page(text),'estimate',['total_amount'],guard_policy='v4')
    assert a['total_amount'] is None and old['guard_version']==v3.GUARD_VERSION
    assert b['total_amount']==1100 and new['guard_version']==GUARD_VERSION


def test_runtime_policy_option(monkeypatch):
    text='금액 안내\n총 청구액: 1,100원'
    monkeypatch.setenv('APP_LLM_GUARD_POLICY','v4')
    monkeypatch.setattr('backend.app.ai.llm_supplement.source_pages',lambda *_:page(text))
    class Client:
        def extract(self,*_):
            return {'ok':True,'proposals':{'total_amount':proposal('1100',text.splitlines()[1],2)}}
    fields,review=supplement_document('example.pdf','',{'total_amount':None},'estimate',client=Client())
    assert fields['total_amount']==1100 and review['guard_version']==GUARD_VERSION


def test_unknown_policy_fails_without_overwriting_rules():
    with pytest.raises(ValueError,match='UNKNOWN_GUARD_POLICY'):
        apply_proposals({}, {'ok':True}, [],'estimate',[],guard_policy='anything')


@pytest.mark.parametrize('kind',['paragraph','table'])
def test_docx_v4_verifies_native_positions_without_claiming_page_layout(tmp_path,kind):
    from docx import Document
    from backend.app.ai.preprocessor import extract_text_from_docx
    document=Document()
    if kind=='paragraph':
        document.add_paragraph('총 청구액: 1,100원')
    else:
        cell=document.add_table(rows=1,cols=1).cell(0,0)
        cell.text='총 청구액\n1,100원'
    path=tmp_path/'native.docx';document.save(path)
    text=extract_text_from_docx(path)
    line=1 if kind=='paragraph' else 2
    d=guard_proposal('total_amount',proposal('1100',text.splitlines()[line-1],line,None),
                     v3.source_pages(path,text),'estimate')
    assert d['accepted']
    for anchor in (d['anchor'],d['anchor']['label_anchor']):
        assert anchor['location_kind']=='docx_structure'
        assert anchor['positions'] and anchor['text_verified']
        assert anchor['page'] is None and anchor['page_layout_verified'] is False


def test_docx_v4_rejects_text_that_has_no_native_location(tmp_path):
    from docx import Document
    from backend.app.ai.preprocessor import extract_text_from_docx
    document=Document();document.add_paragraph('총 청구액: 1,100원')
    path=tmp_path/'native.docx';document.save(path)
    changed=extract_text_from_docx(path).replace('1,100','1,200')
    d=guard_proposal('total_amount',proposal('1200',changed,1,None),
                     v3.source_pages(path,changed),'estimate')
    assert not d['accepted'] and 'DOCX_LOCATION_UNVERIFIED' in d['reasons']


def test_docx_v4_table_cannot_supply_buyer_identity(tmp_path):
    from docx import Document
    from backend.app.ai.preprocessor import extract_text_from_docx
    document=Document();document.add_paragraph('공급자')
    table=document.add_table(rows=2,cols=2)
    table.cell(0,0).text='공급받는자'
    table.cell(1,0).text='상호';table.cell(1,1).text='수신 회사'
    path=tmp_path/'buyer.docx';document.save(path)
    text=extract_text_from_docx(path)
    d=guard_proposal('company_name',proposal('수신 회사','상호\t수신 회사',3,None),
                     v3.source_pages(path,text),'tax_invoice')
    assert not d['accepted'] and 'WRONG_PARTY' in d['reasons']
