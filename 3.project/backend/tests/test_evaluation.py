"""Check metric definitions on tiny hand-authored examples, not dataset answers."""
import copy
import json
from types import SimpleNamespace

import pytest

from backend.evaluation_core import (
    FIELDS, EvaluationNotReady, dataset_file, file_hash, normalize,
    score_predictions, validate_labels, verify_dataset, verify_freeze,
)
from backend.scripts import evaluate
from backend.scripts.import_evaluation_labels import import_workbook
from backend.evaluation_review import blind_check_slots, blank_independent_labels, compare_labels, merge_independent_parts
from backend.scripts.build_reference_labels import design_documents
from backend.scripts.package_independent_review import package
from backend.evaluation_confirmation import apply_human_decisions
from backend.scripts.import_reconciled_labels import import_review


def scoring_example():
    slots = [{"rule_id": "AMOUNT_MATCH_ESTIMATE", "field": "7"},
             {"rule_id": "REQUIRED_DOCUMENTS", "field": "bank_account"}]
    contract = {"contract_id": 1, "split": "dev", "layout_family": "linear",
                "documents": [{"document_id": 7}], "check_slots": slots}
    fields = {field: {"state": "ABSENT", "value": None} for field in FIELDS}
    fields["supply_amount"] = {"state": "VALUE", "value": 100}
    fields["total_amount"] = {"state": "REVIEW", "value": None}
    labels = {"dataset_version": "test-fixture", "split": "dev",
              "review": {"confirmed": True, "reviewer": "test-author", "confirmed_at": "2026-10-01"},
              "documents": {"7": {"reviewed": True, "document_type": "estimate", "fields": fields}},
              "contracts": {"1": {"reviewed": True, "missing_types": ["bank_account"],
                                  "checks": [{**slots[0], "status": "FAIL"}, {**slots[1], "status": "MISSING"}]}}}
    prediction = {"documents": {"7": {"document_type": "estimate", "fields": {}}},
                  "contracts": {"1": {"checks": [{**slots[0], "status": "PASS"}]}}}
    return [contract], labels, prediction


def test_abstained_present_fields_remain_in_accuracy_denominator():
    contracts, labels, prediction = scoring_example()
    result = score_predictions(contracts, labels, prediction)
    assert result["field_counts"]["scored"] == 2
    assert result["field_accuracy"] == 0.5  # Missing a known value is an error.
    assert result["present_field_accuracy"] == 0
    assert result["ambiguous_field_accuracy"] == 1
    assert result["extraction_coverage"] == 0


@pytest.mark.parametrize("proposal", ["invented", "NaN", "-10", 121])
def test_any_nonempty_proposal_on_ambiguous_field_is_an_error(proposal):
    contracts, labels, prediction = scoring_example()
    prediction["documents"]["7"]["fields"] = {"supply_amount": "100원", "total_amount": proposal}
    result = score_predictions(contracts, labels, prediction)
    assert result["field_accuracy"] == 0.5
    assert result["ambiguous_field_accuracy"] == 0
    assert result["extraction_coverage"] == pytest.approx(2 / len(FIELDS))


def test_false_pass_and_missed_missing_document_are_counted():
    contracts, labels, prediction = scoring_example()
    result = score_predictions(contracts, labels, prediction)
    assert result["wrong_pass_rate"] == 0.5  # Two unsafe truths, one falsely passed.
    assert result["wrong_pass_among_passes"] == 1
    assert result["missing_recall"] == 0
    assert result["missing_precision"] is None
    assert result["missing_counts"] == {"tp": 0, "fn": 1, "fp": 0}
    assert len([row for row in result["failures"] if row["kind"] == "check"]) == 2


def test_false_values_on_absent_fields_are_reported_separately():
    contracts, labels, prediction = scoring_example()
    prediction["documents"]["7"]["fields"]["company_name"] = "made-up"
    result = score_predictions(contracts, labels, prediction)
    assert result["false_extraction_on_absent_rate"] == 1 / 5
    assert result["field_accuracy"] == 0.5


def test_duplicate_rule_results_cannot_silently_overwrite_a_failure():
    contracts, labels, prediction = scoring_example()
    prediction["contracts"]["1"]["checks"] *= 2
    with pytest.raises(EvaluationNotReady, match="중복"):
        score_predictions(contracts, labels, prediction)


def test_labels_need_human_review_and_correct_split():
    contracts, labels, _ = scoring_example()
    manifest = {"dataset_version": "test-fixture", "contracts": contracts}
    assert validate_labels(manifest, labels, "dev") == contracts
    unreviewed = copy.deepcopy(labels)
    unreviewed["review"]["confirmed"] = False
    with pytest.raises(EvaluationNotReady, match="사람"):
        validate_labels(manifest, unreviewed, "dev")
    with pytest.raises(EvaluationNotReady, match="버전"):
        validate_labels(manifest, labels, "test")
    labels["contracts"]["1"]["missing_types"] = []
    with pytest.raises(EvaluationNotReady, match="누락 문서 목록"):
        validate_labels(manifest, labels, "dev")


def test_dataset_rejects_changed_documents_and_overlapping_layouts(tmp_path):
    source = tmp_path / "fiction.txt"
    source.write_text("fictional data", encoding="utf-8")
    manifest = {"contracts": [{"contract_id": 1, "split": "dev", "layout_family": "linear",
                              "documents": [{"document_id": 1, "path": source.name, "sha256": file_hash(source)}]}]}
    verify_dataset(tmp_path, manifest)
    changed = copy.deepcopy(manifest)
    changed["contracts"][0]["documents"][0]["sha256"] = "wrong"
    with pytest.raises(EvaluationNotReady, match="변경"):
        verify_dataset(tmp_path, changed)
    test_contract = copy.deepcopy(manifest["contracts"][0])
    test_contract.update(contract_id=2, split="test")
    test_contract["documents"][0]["document_id"] = 2
    manifest["contracts"].append(test_contract)
    with pytest.raises(EvaluationNotReady, match="서식 계열"):
        verify_dataset(tmp_path, manifest)


def test_dataset_paths_cannot_escape_root(tmp_path):
    root = tmp_path / "dataset"
    root.mkdir()
    (tmp_path / "outside.txt").write_text("do not read", encoding="utf-8")
    with pytest.raises(EvaluationNotReady, match="경로"):
        dataset_file(root, "../outside.txt")


def test_freeze_detects_changes_after_review(tmp_path):
    manifest = tmp_path / "manifest.json"
    labels = tmp_path / "ground_truth_dev.json"
    manifest.write_text("{}", encoding="utf-8")
    labels.write_text("{}", encoding="utf-8")
    receipt = {"labels_sha256": file_hash(labels), "manifest_sha256": file_hash(manifest)}
    labels.with_suffix(".freeze.json").write_text(json.dumps(receipt), encoding="utf-8")
    assert verify_freeze(manifest, labels) == receipt
    labels.write_text('{"changed": true}', encoding="utf-8")
    with pytest.raises(EvaluationNotReady, match="변경"):
        verify_freeze(manifest, labels)


def test_unreviewed_labels_block_extraction_and_report_creation(tmp_path, monkeypatch):
    source = tmp_path / "document.txt"
    source.write_text("fiction", encoding="utf-8")
    contracts, labels, _ = scoring_example()
    contracts[0]["documents"][0].update(path=source.name, sha256=file_hash(source))
    labels["review"]["confirmed"] = False
    (tmp_path / "manifest.json").write_text(json.dumps({"dataset_version": "test-fixture", "contracts": contracts}), encoding="utf-8")
    (tmp_path / "ground_truth_dev.json").write_text(json.dumps(labels), encoding="utf-8")
    def forbidden(*args):
        pytest.fail("Unreviewed answers must not run extraction")
    monkeypatch.setattr(evaluate, "run_rule", forbidden)
    output = tmp_path / "reports"
    assert evaluate.main(["--dataset", str(tmp_path), "--output", str(output)]) == 2
    assert not output.exists()


def test_test_split_requires_explicit_final_evaluation(tmp_path):
    assert evaluate.main(["--dataset", str(tmp_path), "--split", "test"]) == 2


def test_number_and_date_normalization_preserves_business_number_zeros():
    assert normalize("business_registration_no", "000-12-34567") == "0001234567"
    assert normalize("total_amount", "1,000원") == 1000
    assert normalize("issue_date", "2026.9.1") == "2026-09-01"
    assert normalize("issue_date", "2026-02-30") is None


def source_example():
    contracts, labels, _ = scoring_example()
    contracts[0]["documents"][0]["path"] = "dev/example.pdf"
    labels["review"] = {"reviewer_type": "source_reference", "confirmed": False,
                        "author": "test-source-author", "created_at": "2026-10-01"}
    for category in ("documents", "contracts"):
        for row in labels[category].values():
            row.update(reviewed=False, source_checked=True)
    return {"dataset_version": "test-fixture", "contracts": contracts}, labels


def independent_example(manifest, reference):
    answer = blank_independent_labels(manifest, manifest["contracts"])
    for key, row in reference["documents"].items():
        answer["documents"][key].update(document_type=row["document_type"], fields=copy.deepcopy(row["fields"]),
                                       original_read={"readable": True, "method": "visual"})
    for key, row in reference["contracts"].items():
        answer["contracts"][key].update(missing_types=row["missing_types"], checks=copy.deepcopy(row["checks"]))
    return answer


def test_source_draft_is_allowed_only_as_unconfirmed_dev_preliminary():
    manifest, labels = source_example()
    assert validate_labels(manifest, labels, "dev", allow_unreviewed=True)
    with pytest.raises(EvaluationNotReady, match="사람"):
        validate_labels(manifest, labels, "dev")
    labels["review"].update(confirmed=True, reviewer="a-model", confirmed_at="2026-10-01")
    with pytest.raises(EvaluationNotReady, match="사람"):
        validate_labels(manifest, labels, "dev")
    with pytest.raises(EvaluationNotReady, match="false"):
        validate_labels(manifest, labels, "dev", allow_unreviewed=True)


def test_preliminary_cannot_unlock_test_split():
    manifest, labels = source_example()
    with pytest.raises(EvaluationNotReady, match="dev"):
        validate_labels(manifest, labels, "test", allow_unreviewed=True)
    assert evaluate.main(["--preliminary", "--split", "test", "--final"]) == 2


def test_design_labels_keep_printed_error_and_abstain_on_unlabeled_numbers():
    contract = {"contract_id": 109, "split": "dev", "scenario": "ambiguous_amount",
                "documents": [{"document_id": index} for index in range(6)]}
    rows = list(design_documents(contract))
    estimate = rows[1][1]
    assert estimate["document_type"] == "estimate"
    assert estimate["fields"]["total_amount"] == {"state": "REVIEW", "value": None}
    contract.update(contract_id=110, scenario="arithmetic_mismatch")
    invoice = list(design_documents(contract))[4][1]["fields"]
    assert invoice["total_amount"]["value"] == invoice["supply_amount"]["value"] + invoice["tax_amount"]["value"] + 1000


def test_blind_slots_do_not_reveal_fixture_document_type():
    contract = {"documents": [{"document_id": 1}, {"document_id": 2}]}
    slots = blind_check_slots(contract)
    first = {row["rule_id"] for row in slots if row["field"] == "1"}
    second = {row["rule_id"] for row in slots if row["field"] == "2"}
    assert first == second
    assert {"AMOUNT_MATCH_ESTIMATE", "AMOUNT_MATCH_TAX_INVOICE"} <= first


def test_comparison_requires_original_reading_and_never_counts_blanks_as_agreement():
    manifest, reference = source_example()
    blank = blank_independent_labels(manifest, manifest["contracts"])
    result = compare_labels(manifest, reference, blank)
    assert result["counts"] == {"UNANSWERED": len(result["rows"])}
    assert result["review"]["independent_complete"] is False
    answer = independent_example(manifest, reference)
    answer["documents"]["7"]["original_read"]["readable"] = False
    assert compare_labels(manifest, reference, answer)["counts"] == result["counts"]


def test_comparison_normalizes_values_and_selects_reproducible_human_queue():
    manifest, reference = source_example()
    answer = independent_example(manifest, reference)
    answer["documents"]["7"]["fields"]["supply_amount"]["value"] = "100원"
    first = compare_labels(manifest, reference, answer, seed=123)
    second = compare_labels(manifest, reference, answer, seed=123)
    assert first == second
    assert first["counts"] == {"AGREE": len(first["rows"])}
    ambiguous = [row for row in first["rows"] if row["key"] == "total_amount"]
    assert ambiguous[0]["selection_reason"] == "AMBIGUITY"
    assert first["sampling"]["sampled_count"] == 2  # ceil(10 nonmandatory agreements * .2)
    assert first["review"]["confirmed"] is False
    assert all(row["human_review_completed"] is False for row in first["rows"])
    answer["documents"]["7"]["fields"]["supply_amount"]["value"] = 120
    mismatch = next(row for row in compare_labels(manifest, reference, answer)["rows"] if row["key"] == "supply_amount")
    assert mismatch["comparison"] == "DISAGREE"
    assert mismatch["human_review_required"] is True


def test_comparison_rejects_unknown_or_duplicate_check_ids():
    manifest, reference = source_example()
    answer = independent_example(manifest, reference)
    answer["contracts"]["1"]["checks"] *= 2
    with pytest.raises(EvaluationNotReady, match="중복"):
        compare_labels(manifest, reference, answer)


def test_independent_merge_rejects_reference_exposure_and_duplicate_batches():
    manifest, reference = source_example()
    answer = independent_example(manifest, reference)
    assert merge_independent_parts([answer])["documents"] == answer["documents"]
    with pytest.raises(EvaluationNotReady, match="중복"):
        merge_independent_parts([answer, answer])
    answer["review"]["used_reference_answers"] = True
    with pytest.raises(EvaluationNotReady, match="미참조"):
        merge_independent_parts([answer])


def confirmation_example():
    manifest, reference = source_example()
    answer = independent_example(manifest, reference)
    comparison = compare_labels(manifest, reference, answer)
    comparison.update(hashes={"reference": "a" * 64, "independent_replies": {"test.json": "b" * 64}}, comparison_sha256="c" * 64)
    decisions = {}
    for row in comparison["rows"]:
        if not row["human_review_required"]:
            continue
        expected = row["reference"]
        state = expected["state"] if isinstance(expected, dict) else "TYPE" if row["key"] == "document_type" else "MISSING_TYPES" if row["key"] == "missing_types" else expected
        value = expected.get("value") if isinstance(expected, dict) else expected if state in {"TYPE", "MISSING_TYPES"} else None
        decisions[row["review_id"]] = {"state": state, "value": value, "confirmation": "확인", "evidence": "hand-authored unit-test decision"}
    return manifest, reference, comparison, decisions


def test_human_interpretation_question_is_required_even_when_both_ai_answers_agree():
    manifest, reference = source_example()
    answer = independent_example(manifest, reference)
    request = {"review_id": "1/7/supply_amount", "reason": "unit-test interpretation question"}
    result = compare_labels(manifest, reference, answer, review_requests=[request])
    row = next(row for row in result["rows"] if row["review_id"] == request["review_id"])
    assert row["comparison"] == "AGREE"
    assert row["human_review_required"] is True
    assert row["selection_reason"] == "INTERPRETATION"
    with pytest.raises(EvaluationNotReady, match="중복"):
        compare_labels(manifest, reference, answer, review_requests=[request, request])


def test_stratified_sample_keeps_total_and_covers_cases_and_document_types():
    import math
    manifest, reference = source_example()
    second = copy.deepcopy(manifest["contracts"][0])
    second.update(contract_id=2)
    second["documents"][0]["document_id"] = 8
    second["check_slots"][0]["field"] = "8"
    manifest["contracts"].append(second)
    reference["documents"]["8"] = copy.deepcopy(reference["documents"]["7"])
    reference["documents"]["8"]["document_type"] = "contract"
    reference["contracts"]["2"] = copy.deepcopy(reference["contracts"]["1"])
    reference["contracts"]["2"]["checks"][0]["field"] = "8"
    independent = independent_example(manifest, reference)
    result = compare_labels(manifest, reference, independent, sample_rate=0.5, sample_method="case_document_type")
    assert result == compare_labels(manifest, reference, independent, sample_rate=0.5, sample_method="case_document_type")
    assert result["sampling"]["sampled_count"] == math.ceil(result["sampling"]["eligible_agreements"] * 0.5)
    assert all(group["sampled_count"] > 0 for group in result["sampling"]["strata"])
    assert {group["contract_id"] for group in result["sampling"]["strata"]} == {"1", "2"}
    assert {group["document_type"] for group in result["sampling"]["strata"]} == {"estimate", "contract", "contract_checks"}


def test_reconciled_labels_preserve_actual_human_coverage_and_ai_agreements():
    manifest, reference, comparison, decisions = confirmation_example()
    labels = apply_human_decisions(manifest, reference, comparison, decisions, "test-human")
    assert validate_labels(manifest, labels, "dev")
    assert labels["review"]["human_full_review"] is False
    assert labels["review"]["human_direct_count"] == len(decisions)
    assert labels["review"]["ai_agreement_count"] == len(comparison["rows"]) - len(decisions)
    assert labels["documents"]["7"]["reviewed"] is False
    assert labels["documents"]["7"]["accepted"] is True
    assert reference["review"]["confirmed"] is False
    broken = copy.deepcopy(labels)
    broken["review"]["human_direct_count"] += 1
    with pytest.raises(EvaluationNotReady, match="집계"):
        validate_labels(manifest, broken, "dev")


@pytest.mark.parametrize("problem", ["missing_row", "unchecked", "no_evidence", "incomplete_independent"])
def test_reconciled_confirmation_blocks_incomplete_human_review(problem):
    manifest, reference, comparison, decisions = confirmation_example()
    first = next(iter(decisions))
    if problem == "missing_row":
        decisions.pop(first)
    elif problem == "unchecked":
        decisions[first]["confirmation"] = None
    elif problem == "no_evidence":
        decisions[first]["evidence"] = ""
    else:
        comparison["review"]["independent_complete"] = False
    with pytest.raises(EvaluationNotReady):
        apply_human_decisions(manifest, reference, comparison, decisions, "test-human")


def test_interpretation_correction_records_human_reason_and_keeps_source_draft():
    manifest, reference = source_example()
    answer = independent_example(manifest, reference)
    comparison = compare_labels(manifest, reference, answer, review_requests=[{"review_id": "1/7/supply_amount", "reason": "hand-authored question"}])
    comparison.update(hashes={"reference": "a" * 64, "independent_replies": {"test.json": "b" * 64}}, comparison_sha256="c" * 64)
    _, _, _, decisions = confirmation_example()
    for row in comparison["rows"]:
        if row["human_review_required"] and row["review_id"] not in decisions:
            expected = row["reference"]
            state = expected["state"] if isinstance(expected, dict) else "TYPE" if row["key"] == "document_type" else "MISSING_TYPES" if row["key"] == "missing_types" else expected
            value = expected.get("value") if isinstance(expected, dict) else expected if state in {"TYPE", "MISSING_TYPES"} else None
            decisions[row["review_id"]] = {"state": state, "value": value, "confirmation": "확인", "evidence": "unit test source correction"}
    decisions = {rid: decision for rid, decision in decisions.items() if any(row["review_id"] == rid and row["human_review_required"] for row in comparison["rows"])}
    decisions["1/7/supply_amount"].update(value=120, evidence="hand-authored unit-test corrected value")
    labels = apply_human_decisions(manifest, reference, comparison, decisions, "test-human")
    assert labels["documents"]["7"]["fields"]["supply_amount"]["value"] == 120
    assert labels["documents"]["7"]["fields"]["supply_amount"]["review"]["method"] == "human_direct"
    assert reference["documents"]["7"]["fields"]["supply_amount"]["value"] == 100


@pytest.mark.parametrize("case", ["partial", "partial_correction", "full_correction", "unchecked", "changed_reply"])
def test_review_import_checks_provenance_and_expands_after_sample_error(tmp_path, monkeypatch, case):
    import openpyxl
    manifest, reference = source_example()
    source = tmp_path / "fiction.txt"
    source.write_text("hand-authored unit-test original", encoding="utf-8")
    manifest["contracts"][0]["documents"][0].update(path=source.name, sha256=file_hash(source))
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    reference_path = tmp_path / "reference_labels_dev.json"
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    reference_path.with_suffix(".freeze.json").write_text(json.dumps({"labels_sha256": file_hash(reference_path), "manifest_sha256": file_hash(manifest_path)}), encoding="utf-8")
    truth_path = tmp_path / "ground_truth_dev.json"
    truth_path.write_text(json.dumps(reference), encoding="utf-8")
    before_hash = file_hash(truth_path)
    reply_dir = tmp_path / "replies"
    reply_dir.mkdir()
    reply = independent_example(manifest, reference)
    reply_path = reply_dir / "test.json"
    reply_path.write_text(json.dumps(reply), encoding="utf-8")
    comparison = compare_labels(manifest, reference, reply, sample_rate=1.0 if case == "full_correction" else 0.20)
    comparison["hashes"] = {"manifest": file_hash(manifest_path), "reference": file_hash(reference_path), "independent_replies": {reply_path.name: file_hash(reply_path)}}
    comparison_path = tmp_path / "comparison.json"
    comparison_path.write_text(json.dumps(comparison), encoding="utf-8")
    workbook = tmp_path / "test-workbook.dat"
    workbook.write_bytes(b"unit-test workbook stand-in")
    rows = {"필드확인": [], "검수확인": []}
    for row in comparison["rows"]:
        if not row["human_review_required"]:
            continue
        expected = row["reference"]
        state = expected["state"] if isinstance(expected, dict) else "TYPE" if row["key"] == "document_type" else "MISSING_TYPES" if row["key"] == "missing_types" else expected
        value = expected.get("value") if isinstance(expected, dict) else json.dumps(expected) if state == "MISSING_TYPES" else expected if state == "TYPE" else None
        if case in {"partial_correction", "full_correction"} and row["key"] == "business_registration_no":
            state, value = "VALUE", "000-12-34567"
        item = [row["review_id"], None, None, None, None, None, None, None, state, value, None if case == "unchecked" else "확인", "unit-test direct source check"]
        rows["필드확인" if row["document_id"] is not None else "검수확인"].append(item)
    class Sheet:
        def __init__(self, rows=(), cells=None):
            self.rows, self.cells = rows, cells or {}
        def iter_rows(self, **kwargs):
            return iter(self.rows)
        def __getitem__(self, cell):
            return SimpleNamespace(value=self.cells[cell])
    class Workbook(dict):
        def close(self):
            pass
    book = Workbook({"안내": Sheet(cells={"B2": manifest["dataset_version"], "B3": file_hash(comparison_path), "B4": file_hash(manifest_path), "B12": "test-human"}),
                     **{name: Sheet(rows=items) for name, items in rows.items()}})
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *args, **kwargs: book)
    if case == "changed_reply":
        reply_path.write_text(json.dumps(reply) + " ", encoding="utf-8")
    if case in {"partial_correction", "unchecked", "changed_reply"}:
        with pytest.raises(EvaluationNotReady, match="표본|직접 확인|Claude"):
            import_review(workbook, tmp_path, comparison_path, reply_dir, "test-human")
        assert file_hash(truth_path) == before_hash
        assert not truth_path.with_suffix(".freeze.json").exists()
    else:
        labels = import_review(workbook, tmp_path, comparison_path, reply_dir, "test-human")
        assert labels["review"]["human_full_review"] is (case == "full_correction")
        assert verify_freeze(manifest_path, truth_path)["workbook_sha256"] == file_hash(workbook)
        if case == "full_correction":
            assert labels["documents"]["7"]["fields"]["business_registration_no"]["value"] == "000-12-34567"


def test_final_report_discloses_human_sample_and_ai_agreement_counts(tmp_path):
    manifest, reference, comparison, decisions = confirmation_example()
    truth = apply_human_decisions(manifest, reference, comparison, decisions, "test-human")
    contracts, _, predictions = scoring_example()
    metrics = score_predictions(contracts, truth, predictions)
    metrics["mean_document_seconds"] = 0.1
    metadata = {"split": "dev", "dataset_version": "test-fixture", "evaluated_at": "unit-test time",
                "labels_sha256": "a" * 64, "manifest_sha256": "b" * 64, "label_review": truth["review"]}
    evaluate.write_report(tmp_path, metadata, metrics)
    text = (tmp_path / "baseline_rule.md").read_text(encoding="utf-8")
    assert f"{truth['review']['human_direct_count']}칸을 직접 확인" in text
    assert f"{truth['review']['ai_agreement_count']}칸은 두 AI 초안의 일치" in text
    assert "사람 전수 검수로 해석하지 않는다" in text


def test_blind_package_contains_dev_originals_but_no_oracle_test_or_predictions(tmp_path):
    import zipfile
    root = tmp_path / "dataset"
    root.mkdir()
    dev = root / "dev.txt"
    dev.write_text("fictional original", encoding="utf-8")
    heldout = root / "test.txt"
    heldout.write_text("held-out original", encoding="utf-8")
    contracts = []
    for cid, split, source in ((1, "dev", dev), (2, "test", heldout)):
        contracts.append({"contract_id": cid, "split": split, "scenario": "must-not-export",
                          "layout_family": split, "description": "oracle description", "input": {"contract_amount": 100},
                          "documents": [{"document_id": cid, "path": source.name, "sha256": file_hash(source)}],
                          "check_slots": [{"rule_id": "TYPE_SPECIFIC_ORACLE", "field": str(cid)}]})
    (root / "manifest.json").write_text(json.dumps({"dataset_version": "synthetic-v1", "created_for": "local-evaluation", "contracts": contracts}), encoding="utf-8")
    (root / "reference_labels_dev.json").write_text("secret answers", encoding="utf-8")
    (root / "rule_predictions.json").write_text("secret predictions", encoding="utf-8")
    archive = package(root, tmp_path / "blind")
    with zipfile.ZipFile(archive) as zipped:
        assert "case_01/dev.txt" in zipped.namelist()
        assert not any("test.txt" in name or "reference_labels" in name or "rule_predictions" in name or "manifest.json" in name for name in zipped.namelist())
        inputs = zipped.read("case_01/registered_contract.json").decode("utf-8")
        assert "scenario" not in inputs and "oracle" not in inputs and "layout_family" not in inputs
        template = zipped.read("case_01/reply_template.json").decode("utf-8")
        assert "TYPE_SPECIFIC_ORACLE" not in template


@pytest.mark.parametrize("wrong_version", [False, True])
def test_human_workbook_import_freezes_answers_and_checks_dataset_version(tmp_path, monkeypatch, wrong_version):
    import openpyxl
    contracts, human, _ = scoring_example()
    source = tmp_path / "fiction.txt"
    source.write_text("fiction", encoding="utf-8")
    contracts[0]["documents"][0].update(path=source.name, sha256=file_hash(source))
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps({"dataset_version": "test-fixture", "contracts": contracts}), encoding="utf-8")
    template = copy.deepcopy(human)
    template["review"]["confirmed"] = False
    for field in template["documents"]["7"]["fields"].values():
        field.update(state=None, value=None)
    for check in template["contracts"]["1"]["checks"]:
        check["status"] = None
    labels_path = tmp_path / "ground_truth_dev.json"
    labels_path.write_text(json.dumps(template), encoding="utf-8")
    workbook_path = tmp_path / "test-input.dat"
    workbook_path.write_bytes(b"test-only workbook stand-in")

    class Sheet:
        def __init__(self, rows=(), cells=None):
            self.rows, self.cells = rows, cells or {}
        def iter_rows(self, **kwargs):
            return iter(self.rows)
        def __getitem__(self, cell):
            return SimpleNamespace(value=self.cells[cell])

    field_rows = [[7, 1, field, "label", answer["state"], answer["value"]]
                  for field, answer in human["documents"]["7"]["fields"].items()]
    book = {
        "안내": Sheet(cells={"B2": "other" if wrong_version else "test-fixture", "B3": file_hash(manifest_path)}),
        "DEV_유형": Sheet(rows=[[7, 1, "fiction.txt", "test", "estimate", "확인"]]),
        "DEV_필드": Sheet(rows=field_rows),
        "DEV_검수": Sheet(rows=[[1, "test", row["rule_id"], row["field"], row["status"], "확인"]
                               for row in human["contracts"]["1"]["checks"]]),
    }
    class Workbook(dict):
        def close(self):
            pass
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *args, **kwargs: Workbook(book))
    if wrong_version:
        with pytest.raises(EvaluationNotReady, match="버전"):
            import_workbook(workbook_path, tmp_path, "dev", "human-reviewer")
        assert not labels_path.with_suffix(".freeze.json").exists()
    else:
        truth = import_workbook(workbook_path, tmp_path, "dev", "human-reviewer")
        assert truth["review"]["confirmed"] is True
        assert truth["review"]["reviewer"] == "human-reviewer"
        assert truth["contracts"]["1"]["missing_types"] == ["bank_account"]
        assert verify_freeze(manifest_path, labels_path)["labels_sha256"] == file_hash(labels_path)
