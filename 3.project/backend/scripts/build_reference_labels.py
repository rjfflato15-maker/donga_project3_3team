"""Derive reference labels from fixture design, never app predictions.

These are source-checked drafts, not human-confirmed ground truth. Explicit
--split test prepares heldout labels without scoring, tuning or predictions.
"""
import argparse
import json
import re
import sys
import tempfile
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))
from backend.evaluation_core import DOCUMENT_TYPES, FIELDS, EvaluationNotReady, file_hash, validate_labels, verify_dataset
from backend.evaluation_review import validate_reference_draft
from backend.scripts.generate_evaluation_dataset import fictional_business_number, make_lines, write_image


def compact(text):
    # Two-column fixtures omit ':' between the same label and value.
    return re.sub(r"[\s:：]", "", text)


def verify_source(path, lines, family, document_id, font_path):
    """Check file content with readers independent of the app extraction path."""
    if path.suffix == ".docx":
        from docx import Document
        document = Document(path)
        chunks = [paragraph.text for paragraph in document.paragraphs]
        chunks += [" ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        if compact("\n".join(chunks)) != compact("\n".join(lines)):
            raise EvaluationNotReady("DOCX와 생성 설계의 내용이 다릅니다.")
        return "independent_docx_content"
    if path.suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if len(reader.pages) != 1:
            raise EvaluationNotReady("합성 PDF의 페이지 수가 설계와 다릅니다.")
        text = reader.pages[0].extract_text()
        if text.strip():
            if compact(text) != compact("\n".join(lines)):
                raise EvaluationNotReady("PDF와 생성 설계의 내용이 다릅니다.")
            return "independent_pdf_content"
        if len(reader.pages[0].images) != 1:
            raise EvaluationNotReady("스캔 PDF의 이미지 수가 설계와 다릅니다.")
        actual = reader.pages[0].images[0].image.convert("RGB")
        noisy = True
        method = "scanned_pdf_design_pixels"
    else:
        from PIL import Image
        with Image.open(path) as source:
            actual = source.convert("RGB")
        noisy = "case_07" in path.as_posix()
        method = "png_design_pixels"
    from PIL import Image, ImageChops
    try:
        with tempfile.TemporaryDirectory(prefix="synthetic-reference-") as scratch:
            expected_path = Path(scratch) / "expected.png"
            write_image(expected_path, lines, family, font_path, noisy=noisy, seed=document_id)
            with Image.open(expected_path) as expected:
                if actual.size != expected.size or ImageChops.difference(actual, expected.convert("RGB")).getbbox() is not None:
                    raise EvaluationNotReady("이미지와 생성 설계의 픽셀이 다릅니다.")
    finally:
        actual.close()
    return method


def design_documents(contract):
    split = contract["split"]
    split_index = 1 if split == "dev" else 2
    index = contract["contract_id"] % 100
    scenario = contract["scenario"]
    vendor = f"가상 {'솔빛' if split == 'dev' else '달빛'}제작소 {index}"
    buyer = f"가상 {'파랑' if split == 'dev' else '초록'}운영사 {index}"
    number = fictional_business_number(contract["contract_id"])
    buyer_no = fictional_business_number(contract["contract_id"] + 1000)
    supply = (index + split_index * 10) * 100000
    tax = supply // 10
    kinds = [kind for kind in DOCUMENT_TYPES if not (scenario == "missing" and kind in {"bank_account", "inspection_confirmation"})]
    if scenario == "installments":
        kinds.append("tax_invoice")
    if len(kinds) != len(contract["documents"]):
        raise EvaluationNotReady("설계와 실제 문서 목록의 길이가 다릅니다.")
    invoice_number = 0
    for document, kind in zip(contract["documents"], kinds):
        actual_supply, actual_tax, actual_total = supply, tax, supply + tax
        business = fictional_business_number(contract["contract_id"] + 2000) if scenario == "wrong_business_no" and kind == "estimate" else number
        if scenario == "amount_mismatch" and kind == "estimate":
            actual_supply += 100000
            actual_tax = actual_supply // 10
            actual_total = actual_supply + actual_tax
        partial = scenario == "installments" and kind == "tax_invoice"
        if partial:
            invoice_number += 1
            actual_supply = supply * (3 if invoice_number == 1 else 7) // 10
            actual_tax = actual_supply // 10
            actual_total = actual_supply + actual_tax
        if scenario == "arithmetic_mismatch" and kind == "tax_invoice":
            actual_total += 1000
        day = f"2026-09-{index:02d}"
        lines = make_lines(kind, vendor, buyer, business, buyer_no, actual_supply, actual_tax, actual_total, day, scenario, partial)
        fields = {field: {"state": "ABSENT", "value": None} for field in FIELDS}
        for field, value in (("company_name", vendor), ("business_registration_no", business), ("issue_date", day)):
            fields[field] = {"state": "VALUE", "value": value}
        monetary = kind in {"contract", "estimate", "tax_invoice", "inspection_confirmation"}
        ambiguous = scenario == "ambiguous_amount" and kind == "estimate"
        if ambiguous:
            # Reference numbers have no supply/tax/total meaning, even though
            # the generator knows how it computed them.
            for field in ("supply_amount", "tax_amount", "total_amount", "amount_scope"):
                fields[field] = {"state": "REVIEW", "value": None}
        elif monetary:
            for field, value in (("supply_amount", actual_supply), ("tax_amount", actual_tax), ("total_amount", actual_total)):
                fields[field] = {"state": "VALUE", "value": value}
            fields["amount_scope"] = {"state": "VALUE", "value": "PARTIAL" if partial else "FULL"}
        yield document, {"document_type": kind, "fields": fields}, lines


def reference_checks(contract, documents):
    """Label the published specification, not the current RuleEngine outputs."""
    answers = {(slot["rule_id"], slot["field"]): "N/A" for slot in contract["check_slots"]}
    kinds = {row["document_type"] for row in documents.values()}
    missing = sorted(set(DOCUMENT_TYPES) - kinds)
    for kind in missing:
        answers[("REQUIRED_DOCUMENTS", kind)] = "MISSING"
    answers[("BUSINESS_STATUS", "business_status")] = "REVIEW"
    inputs = contract["input"]
    expected = inputs["contract_amount"]
    if inputs["amount_policy"]["amount_basis"] == "SUPPLY":
        expected += inputs["amount_policy"]["tax_amount"]
    invoices = []
    for did, row in documents.items():
        kind, fields = row["document_type"], row["fields"]
        if kind in {"contract", "estimate", "tax_invoice", "business_registration"}:
            matches = compact(fields["business_registration_no"]["value"]).replace("-", "") == compact(inputs["business_number"]).replace("-", "")
            answers[("BUSINESS_NO_MATCH", did)] = "PASS" if matches else "FAIL"
        if kind not in {"estimate", "tax_invoice"}:
            continue
        if fields["total_amount"]["state"] != "VALUE":
            answers[("AMOUNT_EXTRACTION", did)] = "REVIEW"
            continue
        total = fields["total_amount"]["value"]
        if total != fields["supply_amount"]["value"] + fields["tax_amount"]["value"]:
            answers[("AMOUNT_COMPONENTS", did)] = "FAIL"
        if kind == "tax_invoice" and inputs["amount_policy"]["payment_mode"] == "INSTALLMENT":
            invoices.append(total)
        else:
            answers[("AMOUNT_MATCH_" + kind.upper(), did)] = "PASS" if total == expected else "FAIL"
    if inputs["amount_policy"]["payment_mode"] == "INSTALLMENT":
        total = sum(invoices)
        answers[("INSTALLMENT_TOTAL", "tax_invoice")] = "PASS" if total == expected else ("FAIL" if total > expected else "REVIEW")
    return missing, [{**slot, "status": answers[(slot["rule_id"], slot["field"])]} for slot in contract["check_slots"]]


def build(root, font_path, output=None, *, split="dev"):
    manifest_path = root / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("dataset_version") != "synthetic-v1" or manifest.get("created_for") != "local-evaluation":
        raise EvaluationNotReady("이 생성 정보는 synthetic-v1 가상 자료에만 적용됩니다.")
    verify_dataset(root, manifest, split=split)
    output = output or root / f"reference_labels_{split}.json"
    if output.exists() or output.with_suffix(".freeze.json").exists():
        raise EvaluationNotReady("기존 정답 초안과 버전 기록을 덮어쓰지 않습니다.")
    labels = {"dataset_version": manifest["dataset_version"], "split": split, "reference_version": "source-reference-v1",
              "review": {"confirmed": False, "reviewer_type": "source_reference", "author": "Codex",
                         "created_at": datetime.now(timezone(timedelta(hours=9))).isoformat(),
                         "human_reviewed": False, "claude_reviewed": False,
                         "docx_page_visual_verified": False}, "documents": {}, "contracts": {}}
    methods = Counter()
    for contract in (row for row in manifest["contracts"] if row["split"] == split):
        rows = {}
        for document, answer, lines in design_documents(contract):
            method = verify_source(root / document["path"], lines, contract["layout_family"], document["document_id"], font_path)
            methods[method] += 1
            did = str(document["document_id"])
            rows[did] = answer
            labels["documents"][did] = {**answer, "reviewed": False, "source_checked": True,
                                         "source_check_method": method, "source_sha256": document["sha256"]}
        missing, checks = reference_checks(contract, rows)
        labels["contracts"][str(contract["contract_id"])] = {"reviewed": False, "source_checked": True, "missing_types": missing, "checks": checks}
    validate_reference_draft(manifest, labels, split)
    output.write_text(json.dumps(labels, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    source_files = {path: file_hash(PROJECT / path) for path in (
        "backend/scripts/generate_evaluation_dataset.py", "backend/scripts/build_reference_labels.py", "eval/LABELING_GUIDE.md")}
    receipt = {"labels_sha256": file_hash(output), "manifest_sha256": file_hash(manifest_path),
               "source_files": source_files, "source_check_counts": dict(methods), **labels["review"]}
    output.with_suffix(".freeze.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return labels


def main(argv=None):
    parser = argparse.ArgumentParser(description="합성 생성 정보에서 dev 정답 초안 작성")
    parser.add_argument("--dataset", type=Path, default=PROJECT / "eval/dataset")
    parser.add_argument("--split", choices=["dev", "test"], default="dev")
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/malgun.ttf"))
    args = parser.parse_args(argv)
    try:
        labels = build(args.dataset.resolve(), args.font, split=args.split)
        print(f"{args.split} 생성 정보 정답 초안: 계약 {len(labels['contracts'])}건, 문서 {len(labels['documents'])}개. 사람과 Claude 검수는 미완료입니다.")
        return 0
    except (EvaluationNotReady, OSError, ValueError) as exc:
        print("정답 초안 생성 중단: " + str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
