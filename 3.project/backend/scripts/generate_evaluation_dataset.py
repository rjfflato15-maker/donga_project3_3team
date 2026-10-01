"""Create fictional documents and BLANK labels, never model-derived answers."""
import argparse
import json
import random
import sys
from datetime import datetime
from pathlib import Path

PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))
from backend.evaluation_core import DOCUMENT_TYPES, FIELDS, file_hash

SCENARIOS = (
    ("normal", "6종 제출 및 금액과 번호 일치"),
    ("missing", "통장사본과 검수확인서 미제출"),
    ("wrong_business_no", "견적서의 공급자 사업자번호가 계약 등록 기준과 다름"),
    ("amount_mismatch", "견적 합계가 계약 등록 기준과 다름"),
    ("vat_separate", "공급가액 기준 계약과 세액을 더한 증빙 비교"),
    ("installments", "선금과 잔금 세금계산서 두 건 및 누적 합계"),
    ("scan_noise", "텍스트 없는 PDF와 회전 및 노이즈가 있는 이미지"),
    ("amount_table", "품목 표와 공급가액 세액 합계"),
    ("ambiguous_amount", "의미와 적용 범위를 확정할 수 없는 참고 금액 두 개"),
    ("arithmetic_mismatch", "표시 합계가 공급가액과 세액의 합과 다름"),
)
TITLES = {
    "contract": "외주용역 계약서", "estimate": "견적서",
    "business_registration": "사업자등록증", "bank_account": "통장사본",
    "tax_invoice": "전자세금계산서", "inspection_confirmation": "납품 검수확인서",
}


def fictional_business_number(serial):
    # Prefix 000 is for test material only. A checksum is not proof of registration.
    first = "000" + f"{serial:06d}"[-6:]
    digits = list(map(int, first))
    total = sum(digit * weight for digit, weight in zip(digits, (1, 3, 7, 1, 3, 7, 1, 3, 5))) + digits[8] * 5 // 10
    number = first + str((10 - total % 10) % 10)
    return number[:3] + "-" + number[3:5] + "-" + number[5:]


def make_lines(kind, vendor, buyer, business_no, buyer_no, supply, tax, total, day, scenario, part=False):
    rows = [TITLES[kind], "평가용 가상 서류이며 실제 거래에 사용할 수 없습니다."]
    if kind == "contract":
        rows += ["계약명: 가상 업무지원 시스템 구축", "발주자: " + buyer,
                 "수급인(을): " + vendor, "사업자등록번호: " + business_no]
    elif kind in {"estimate", "tax_invoice"}:
        rows += ["공급자", "상호: " + vendor, "사업자등록번호: " + business_no,
                 "공급받는자", "상호: " + buyer, "사업자등록번호: " + buyer_no]
    elif kind == "business_registration":
        rows += ["상호: " + vendor, "사업자등록번호: " + business_no,
                 "대표자: 가상 담당자", "업태: 서비스업", "종목: 소프트웨어 개발"]
    elif kind == "bank_account":
        rows += ["예금주: " + vendor, "은행명: 가상은행", "계좌번호: TEST-ACCOUNT-001",
                 "사업자등록번호: " + business_no]
    else:
        rows += ["계약상대자: " + vendor, "사업자등록번호: " + business_no,
                 "납품 내역: 업무지원 시스템", "검수 결과: 납품 항목 확인"]
    if kind in {"contract", "estimate", "tax_invoice", "inspection_confirmation"}:
        if scenario == "ambiguous_amount" and kind == "estimate":
            rows += [f"참고값 A: {supply:,}원", f"참고값 B: {total:,}원", "금액의 적용 범위가 표시되지 않음"]
        else:
            if scenario == "amount_table":
                rows += ["품목 | 수량 | 단가 | 금액", f"시스템 개발 | 1 | {supply:,} | {supply:,}"]
            rows += [f"공급가액: {supply:,}원", f"세액: {tax:,}원", f"합계금액: {total:,}원"]
        if kind == "tax_invoice":
            rows += ["분할 청구: 선금 또는 잔금" if part else "전액 청구: 전체 계약 대금"]
    rows += ["발행일: " + day]
    return rows


def write_pdf(path, lines, family, font_name):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.pagesizes import A4
    canvas = Canvas(str(path), pagesize=A4, invariant=1)
    canvas.setTitle(lines[0])
    canvas.setAuthor("Synthetic evaluation")
    canvas.setFont(font_name, 17)
    canvas.drawString(48, 790, lines[0])
    y = 758
    canvas.setFont(font_name, 10)
    for line in lines[1:]:
        if family == "test_matrix" and ": " in line:
            key, value = line.split(": ", 1)
            canvas.drawString(48, y, key)
            canvas.drawString(210, y, value)
        else:
            canvas.drawString(48, y, line)
        y -= 22 if family == "test_matrix" else 19
    canvas.showPage()
    canvas.save()


def write_image(path, lines, family, font_path, noisy=False, seed=0):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path), 34)
    body_font = ImageFont.truetype(str(font_path), 25)
    draw.text((95, 95), lines[0], fill="black", font=title_font)
    y = 175
    for line in lines[1:]:
        if family == "test_matrix" and ": " in line:
            key, value = line.split(": ", 1)
            draw.text((95, y), key, fill="black", font=body_font)
            draw.text((440, y), value, fill="black", font=body_font)
        else:
            draw.text((95, y), line, fill="black", font=body_font)
        y += 43
    if noisy:
        rng = random.Random(seed)
        for _ in range(180):
            x, y = rng.randrange(image.width), rng.randrange(image.height)
            draw.point((x, y), fill=(200, 200, 200))
        image = image.rotate(0.65, expand=False, fillcolor="white")
    image.save(path)
    image.close()


def write_docx(path, lines, family):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml.ns import qn
    document = Document()
    document.core_properties.author = "Synthetic evaluation"
    document.core_properties.created = datetime(2026, 10, 1)
    document.core_properties.modified = datetime(2026, 10, 1)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.7)
    for style_name in ("Normal", "Title"):
        style = document.styles[style_name]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(11 if style_name == "Normal" else 17)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")
    document.add_paragraph(lines[0], "Title")
    document.add_paragraph(lines[1])
    if family == "test_matrix":
        table = document.add_table(rows=0, cols=2)
        for line in lines[2:]:
            cells = table.add_row().cells
            if ": " in line:
                cells[0].text, cells[1].text = line.split(": ", 1)
            else:
                cells[0].text = line
                cells[1].text = ""
    else:
        for line in lines[2:]:
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(5)
    document.save(path)


def check_slots(documents, planned_kinds):
    slots = []
    def add(rule, field):
        entry = {"rule_id": rule, "field": str(field)}
        if entry not in slots:
            slots.append(entry)
    for dtype in DOCUMENT_TYPES:
        add("REQUIRED_DOCUMENTS", dtype)
    add("BUSINESS_NO_MATCH", "business_registration_no")
    add("BUSINESS_STATUS", "business_status")
    add("AMOUNT_BASIS", "amount_basis")
    for document, kind in zip(documents, planned_kinds):
        did = document["document_id"]
        for rule in ("ORIGINAL_FILE", "DOCUMENT_TYPE", "DOCUMENT_PROCESSING"):
            add(rule, did)
        if kind in {"contract", "estimate", "tax_invoice", "business_registration"}:
            add("BUSINESS_NO_MATCH", did)
        if kind in {"estimate", "tax_invoice"}:
            for rule in ("AMOUNT_EXTRACTION", "AMOUNT_COMPONENTS", "AMOUNT_MATCH_" + kind.upper()):
                add(rule, did)
        if kind == "tax_invoice":
            add("PAYMENT_SCOPE", did)
    for rule in ("INSTALLMENT_TOTAL", "PAYMENT_SCOPE", "DUPLICATE_INVOICE"):
        add(rule, "tax_invoice")
    return slots


def generate(root, font_path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.pagesizes import A4
    if root.exists() and any(root.iterdir()):
        raise ValueError("평가 자료 폴더가 비어 있지 않습니다. 기존 서류와 정답을 덮어쓰지 않았습니다.")
    if not font_path.is_file():
        raise ValueError("한글 폰트 파일을 지정하세요.")
    pdfmetrics.registerFont(TTFont("EvalKorean", str(font_path)))
    manifest = {"dataset_version": "synthetic-v1", "created_for": "local-evaluation", "contracts": []}
    labels = {split: {"dataset_version": manifest["dataset_version"], "split": split,
                      "review": {"confirmed": False, "reviewer": None, "confirmed_at": None},
                      "documents": {}, "contracts": {}} for split in ("dev", "test")}
    for split_index, split in enumerate(("dev", "test"), 1):
        family = "dev_linear" if split == "dev" else "test_matrix"
        for index, (scenario, description) in enumerate(SCENARIOS, 1):
            cid = split_index * 100 + index
            vendor = f"가상 {'솔빛' if split == 'dev' else '달빛'}제작소 {index}"
            buyer = f"가상 {'파랑' if split == 'dev' else '초록'}운영사 {index}"
            number = fictional_business_number(cid)
            other_no = fictional_business_number(cid + 1000)
            supplier_alt = fictional_business_number(cid + 2000)
            supply = (index + split_index * 10) * 100000
            tax, total = supply // 10, supply + supply // 10
            folder = root / split / f"case_{index:02d}"
            folder.mkdir(parents=True)
            contract = {"contract_id": cid, "split": split, "scenario": scenario,
                        "description": description, "layout_family": family,
                        "input": {"title": "가상 업무지원 시스템 구축", "vendor_name": vendor,
                                  "business_number": number,
                                  "contract_amount": supply if scenario == "vat_separate" else total,
                                  "amount_policy": {"amount_basis": "SUPPLY" if scenario == "vat_separate" else "TOTAL",
                                                    "tax_amount": tax if scenario == "vat_separate" else None,
                                                    "payment_mode": "INSTALLMENT" if scenario == "installments" else "FULL"}},
                        "documents": []}
            kinds = [kind for kind in DOCUMENT_TYPES if not (scenario == "missing" and kind in {"bank_account", "inspection_confirmation"})]
            if scenario == "installments":
                kinds.append("tax_invoice")
            invoice_number = 0
            for doc_index, kind in enumerate(kinds, 1):
                did = cid * 100 + doc_index
                actual_supply, actual_tax, actual_total = supply, tax, total
                business = supplier_alt if scenario == "wrong_business_no" and kind == "estimate" else number
                if scenario == "amount_mismatch" and kind == "estimate":
                    actual_supply += 100000
                    actual_tax, actual_total = actual_supply // 10, actual_supply + actual_supply // 10
                part = scenario == "installments" and kind == "tax_invoice"
                if part:
                    invoice_number += 1
                    actual_supply = supply * (3 if invoice_number == 1 else 7) // 10
                    actual_tax, actual_total = actual_supply // 10, actual_supply + actual_supply // 10
                if scenario == "arithmetic_mismatch" and kind == "tax_invoice":
                    actual_total += 1000
                lines = make_lines(kind, vendor, buyer, business, other_no, actual_supply, actual_tax, actual_total,
                                   f"2026-09-{index:02d}", scenario, part)
                extension = ("pdf", "png", "docx")[(doc_index + index + split_index) % 3]
                scanned = scenario == "scan_noise"
                if scanned:
                    extension = "pdf" if doc_index == 1 else "png"
                path = folder / f"document_{doc_index:02d}.{extension}"
                if extension == "pdf" and scanned:
                    scratch = folder / ".scan-page.png"
                    write_image(scratch, lines, family, font_path, noisy=True, seed=did)
                    canvas = Canvas(str(path), pagesize=A4, invariant=1)
                    canvas.drawImage(str(scratch), 0, 0, width=A4[0], height=A4[1])
                    canvas.save()
                    scratch.unlink()
                elif extension == "pdf":
                    write_pdf(path, lines, family, "EvalKorean")
                elif extension == "png":
                    write_image(path, lines, family, font_path, noisy=scanned, seed=did)
                else:
                    write_docx(path, lines, family)
                document = {"document_id": did, "path": path.relative_to(root).as_posix(), "sha256": file_hash(path)}
                contract["documents"].append(document)
                labels[split]["documents"][str(did)] = {"document_type": None, "reviewed": False,
                    "fields": {field: {"state": None, "value": None} for field in FIELDS}}
            contract["check_slots"] = check_slots(contract["documents"], kinds)
            labels[split]["contracts"][str(cid)] = {"reviewed": False, "missing_types": None,
                "checks": [{**slot, "status": None} for slot in contract["check_slots"]]}
            manifest["contracts"].append(contract)
    (root / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    for split, truth in labels.items():
        (root / f"ground_truth_{split}.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=PROJECT / "eval" / "dataset")
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/malgun.ttf"))
    args = parser.parse_args()
    try:
        manifest = generate(args.output.resolve(), args.font)
    except (ValueError, OSError) as exc:
        print("자료 생성 중단: " + str(exc), file=sys.stderr)
        return 2
    documents = [document for contract in manifest["contracts"] for document in contract["documents"]]
    print(f"가상 자료 생성: 계약 {len(manifest['contracts'])}건, 문서 {len(documents)}개. 정답은 모두 미작성 상태입니다.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
