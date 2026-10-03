"""Prepare hard-v2 review; never run OCR, inference, a guard or final scoring."""
import argparse
import json
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path

PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))
from backend.evaluation_core import EvaluationNotReady, file_hash, verify_dataset, verify_freeze
from backend.evaluation_review import compare_labels, merge_independent_parts, validate_reference_draft
from backend.prediction_artifacts import write_json
from backend.scripts.compare_reference_labels import main as compare_main
from backend.scripts.prepare_test_review import page_builder, review_questions, workbook_from_comparison
from backend.scripts.import_reconciled_labels import import_review, read_wording_resolutions


def paths(project):
    return {'dataset':project/'eval/dataset-v2', 'blind':project/'eval/claude_blind_review_test2',
            'page':project/'eval/human_review_page_test2', 'workbook':project/'eval/human_review_test2.xlsx',
            'comparison':project/'eval/reconciliation-test2-v1',
            'questions':project/'eval/human_questions_test2_v1.json',
            'human':project/'eval/human_reviews/test2', 'output':project/'docs/eval/test2/preparation'}


def policy_hashes(project):
    sources = [*project.glob('backend/app/ai/*.py'), *project.glob('backend/app/validation/*.py')]
    sources += [project/name for name in ('backend/app/schemas/document.py','backend/app/schemas/common.py',
        'backend/app/schemas/validation.py','backend/app/config.py','backend/app/integrations/nts_client.py',
        'backend/evaluation_core.py','backend/scripts/compare_llm_modes.py')]
    return {path.relative_to(project).as_posix():sha256(path.read_bytes().replace(b'\r\n',b'\n')).hexdigest()
            for path in sorted(sources)}


def freeze_policy(project):
    """No heldout originals or answer content are read here."""
    p = paths(project)
    dev = project/'docs/eval/dev2/guard_v4'
    scores = json.loads((dev/'preliminary_comparison.json').read_text(encoding='utf-8'))
    receipt = json.loads((dev/'input_receipt.json').read_text(encoding='utf-8'))
    hashes = policy_hashes(project)
    if (scores['metadata']['v4_sha256'] != hashes['backend/app/ai/grounding_v4.py']
            or receipt['model']['name'] != 'qwen3:latest'
            or receipt['prompt_version'] != 'grounded-fields-v3'):
        raise EvaluationNotReady('dev2에서 비교한 가드·모델·프롬프트와 다릅니다.')
    lock = p['output']/'release_lock.json'
    definition = {'release_version':'guard-v4-test2-fixed', 'frozen':True,
        'guard_policy':'v4', 'guard_version':'grounding-v4-dev2', 'prompt_version':receipt['prompt_version'],
        'selected_model':receipt['model']['name'], 'model_digest':receipt['model']['digest'],
        'settings':receipt['settings'], 'nts_mode':'disabled', 'implementation_sha256':hashes,
        'code_hash_policy':'SHA256 with CRLF normalized to LF',
        'dev2_comparison_sha256':file_hash(dev/'preliminary_comparison.json'),
        'manifest_sha256':file_hash(p['dataset']/'manifest.json'),
        'test_policy':'Independent review and human confirmation, then one reserved final inference; no test2 tuning.',
        'disclosure':'v3 test 결과를 본 뒤 v4 개발을 선택했으며 개발 근거는 dev2만 사용했다.',
        'test2_sources_used_for_tuning':False, 'final_inference_started':False}
    if lock.exists():
        saved = json.loads(lock.read_text(encoding='utf-8'))
        if {k:saved.get(k) for k in definition} != definition:
            raise EvaluationNotReady('고정 이후 구현·자료·개발 기록이 변경되었습니다. test2로 조정하지 않습니다.')
    else:
        saved = {**definition,'frozen_at':datetime.now(timezone.utc).isoformat()}
        write_json(lock,saved,exclusive=True)
    return saved


def verify_blind(p):
    record = json.loads((p['blind']/'package_receipt.json').read_text(encoding='utf-8'))
    if (record.get('dataset_version') != 'synthetic-hard-v2' or record.get('split') != 'test'
            or record.get('contains_reference_answers') is not False
            or record.get('contains_app_predictions') is not False):
        raise EvaluationNotReady('test2 독립 판독 묶음의 출처가 다릅니다.')
    for name,digest in record['files'].items():
        source = (p['blind']/name).resolve()
        if not source.is_relative_to(p['blind'].resolve()) or file_hash(source)!=digest:
            raise EvaluationNotReady('독립 판독 묶음 파일이 수령 해시와 다릅니다.')
    with zipfile.ZipFile(p['blind'].with_suffix('.zip')) as archive:
        expected = {*record['files'], 'package_receipt.json'}
        if set(archive.namelist()) != expected:
            raise EvaluationNotReady('test2 ZIP 파일 목록이 판독 폴더와 다릅니다.')
        for name in expected:
            if sha256(archive.read(name)).hexdigest()!=file_hash(p['blind']/name):
                raise EvaluationNotReady('test2 ZIP 내용이 판독 폴더와 다릅니다.')
    return record


def update_blind_prompt(project):
    """Correct only review instructions before replies; preserve original PDFs."""
    p = paths(project)
    freeze_policy(project)
    receipt = verify_blind(p)
    path = p['blind']/'CLAUDE_PROMPT.md'
    prompt = path.read_text(encoding='utf-8')
    if 'test2 case_01~06' in prompt:
        return
    if any((p['blind']/'replies').glob('*.json')):
        raise EvaluationNotReady('판독을 시작한 묶음의 프롬프트를 변경하지 않습니다.')
    before = {'prompt_sha256':file_hash(path),'zip_sha256':file_hash(p['blind'].with_suffix('.zip')),
              'receipt_sha256':file_hash(p['blind']/'package_receipt.json')}
    prompt = prompt.replace('eval/claude_blind_review 폴더','eval/claude_blind_review_test2 폴더')
    prompt += ('\n## test2 case_01~06 전체 판독\n\n'
        '로컬 작업 범위는 C:\\3조 색인\\3.project\\eval\\claude_blind_review_test2 폴더 내부입니다. '
        '이 묶음 밖의 자료·생성 코드·정답 초안·앱 코드·성능 보고서는 열지 마세요. '
        '다른 작업 브랜치에 복사한 예전 자료와 섞지 마세요.\n\n'
        'case_01부터 case_06까지 6개 계약·36개 PDF를 한 번에 판독하세요. '
        '계약마다 replies/case_01.json부터 replies/case_06.json까지 실제 판독 답을 저장하세요. '
        '웹에서는 해당 이름의 JSON 파일로 반환하세요. 다른 파일은 수정하지 말고 커밋·PR은 하지 마세요. '
        '정답을 모르거나 읽지 못하면 그 사실을 표시하고 추측하지 마세요. '
        '마지막에 원문 확인 건수와 해석 질문만 간단히 알려 주세요.\n\n'
        '내부 split은 test, dataset_version은 synthetic-hard-v2입니다. '
        '템플릿 ID와 키를 유지하고 confirmed=false를 바꾸지 마세요. '
        '사람 확인과 최종 성능 평가는 아직 하지 않은 상태입니다.\n')
    path.write_text(prompt,encoding='utf-8',newline='\n')
    receipt['files']['CLAUDE_PROMPT.md']=file_hash(path)
    receipt.update(instruction_version='test2-blind-v2',parent_receipt_sha256=before['receipt_sha256'])
    write_json(p['blind']/'package_receipt.json',receipt)
    archive_path=p['blind'].with_suffix('.zip')
    temporary=archive_path.with_suffix('.zip.tmp')
    try:
        with zipfile.ZipFile(temporary,'w',compression=zipfile.ZIP_DEFLATED) as archive:
            for name in (*sorted(receipt['files']),'package_receipt.json'):
                archive.write(p['blind']/name,name)
        temporary.replace(archive_path)
    finally:
        temporary.unlink(missing_ok=True)
    verify_blind(p)
    write_json(p['output']/'blind_prompt_update.json',{'before':before,
        'after':{'prompt_sha256':file_hash(path),'zip_sha256':file_hash(archive_path),
                 'receipt_sha256':file_hash(p['blind']/'package_receipt.json')},
        'original_pdf_hashes_unchanged':True,'reference_answers_unchanged':True})


def readiness(project):
    lock = freeze_policy(project)
    p = paths(project)
    receipt = verify_blind(p)
    files = sorted((p['blind']/'replies').glob('*.json'))
    workbook_ready = p['workbook'].is_file()
    comparison_path = p['comparison']/'comparison.json'
    independent_complete = (json.loads(comparison_path.read_text(encoding='utf-8'))['review']['independent_complete']
                            if comparison_path.exists() else False)
    labels = json.loads((p['dataset']/'ground_truth_test.json').read_text(encoding='utf-8'))
    confirmed = labels.get('review',{}).get('confirmed') is True
    return {'dataset':'synthetic-hard-v2','corpus':'test2','contracts':6,'documents':36,
        'blind_package_verified_files':len(receipt['files']), 'independent_reply_files':len(files),
        'independent_review_complete':independent_complete,'human_page_ready':workbook_ready,
        'human_labels_confirmed':confirmed, 'final_inference_started':(p['dataset']/'final_evaluation_attempt.json').exists(),
        'status':'labels_confirmed' if confirmed else 'human_review_pending' if workbook_ready else 'independent_review_pending',
        'release_sha256':file_hash(p['output']/'release_lock.json'),
        'model':lock['selected_model'],'prompt_version':lock['prompt_version'],'guard_policy':lock['guard_policy']}


def reconcile(project, replies):
    freeze_policy(project)
    p = paths(project)
    verify_blind(p)
    if replies.resolve() != (p['blind']/'replies').resolve():
        raise EvaluationNotReady('test2 전용 replies 폴더만 사용하세요.')
    files = sorted(replies.glob('*.json'))
    if {path.name for path in files} != {f'case_{i:02d}.json' for i in range(1,7)}:
        raise EvaluationNotReady('Claude의 test2 case_01~06 실제 판독 JSON 6개가 먼저 필요합니다.')
    if p['workbook'].exists() or p['comparison'].exists():
        raise EvaluationNotReady('이미 만든 사람 확인 목록을 덮어쓰지 않습니다.')
    root = p['dataset']
    manifest = json.loads((root/'manifest.json').read_text(encoding='utf-8'))
    reference = json.loads((root/'reference_labels_test.json').read_text(encoding='utf-8'))
    verify_dataset(root,manifest,split='test')
    verify_freeze(root/'manifest.json',root/'reference_labels_test.json')
    validate_reference_draft(manifest,reference,'test')
    independent = merge_independent_parts([json.loads(f.read_text(encoding='utf-8-sig')) for f in files],split='test')
    if not compare_labels(manifest,reference,independent)['review']['independent_complete']:
        raise EvaluationNotReady('독립 판독에 미응답 또는 원문 미판독이 있습니다. 먼저 보완하세요.')
    questions = review_questions(reference)
    for did,doc in reference['documents'].items():
        for field,answer in doc['fields'].items():
            if answer['state']=='REVIEW':
                questions.append({'review_id':f'{int(did)//100}/{did}/{field}',
                    'reason':'원문에서 가려지거나 뜻을 확정할 수 없는 값은 REVIEW입니다. 등록 금액이나 다른 숫자의 합으로 복원하지 않습니다.'})
    if p['questions'].exists():
        if json.loads(p['questions'].read_text(encoding='utf-8'))!=questions:
            raise EvaluationNotReady('이미 만든 해석 질문을 덮어쓰지 않습니다.')
    else:
        write_json(p['questions'],questions,exclusive=True)
    args = [*[str(f) for f in files],'--dataset',str(root),'--split','test',
        '--output',str(p['comparison']),'--sample-rate','0.10','--sample-method','case_document_type',
        '--human-questions',str(p['questions'])]
    if compare_main(args)!=0:
        raise EvaluationNotReady('test2 대조에 실패했습니다.')
    workbook = workbook_from_comparison(root,p['comparison']/'comparison.json',p['workbook'])
    page_builder().build_page(split='test',dataset=root,blind=p['blind'],output=p['page'],
                              source_directory=root/'test2',workbook=workbook)
    html_path = p['page']/'index.html'
    html = html_path.read_text(encoding='utf-8')
    html_path.write_text(html.replace('human_review_result_${DATA.split||"dev"}', 'human_review_result_test2'),
                         encoding='utf-8',newline='\n')
    comparison = json.loads((p['comparison']/'comparison.json').read_text(encoding='utf-8'))
    return {'human_page_ready':True,'human_review_required':comparison['required_human_count'],
        'counts':comparison['counts'], 'selection':dict(Counter(r['selection_reason'] for r in comparison['rows'] if r['human_review_required'])),
        'final_inference_started':False}


def confirm(project, result_path, resolution_path=None):
    import openpyxl
    freeze_policy(project)
    p = paths(project)
    if not p['workbook'].exists():
        raise EvaluationNotReady('독립 판독 대조와 사람 확인 목록을 먼저 완료하세요.')
    result = json.loads(result_path.read_text(encoding='utf-8-sig'))
    reviewer = str(result.get('reviewer','')).strip()
    if result.get('split')!='test' or result.get('source_xlsx_sha256')!=file_hash(p['workbook']) or not reviewer:
        raise EvaluationNotReady('검수자·test2 원본 양식 해시가 다릅니다. 기존 test 결과를 재사용하지 않습니다.')
    if p['human'].exists():
        raise EvaluationNotReady('기존 사람 확인 출처를 덮어쓰지 않습니다.')
    manifest = json.loads((p['dataset']/'manifest.json').read_text(encoding='utf-8'))
    wording_resolutions = read_wording_resolutions(resolution_path, result_path, manifest)
    book = openpyxl.load_workbook(p['workbook'])
    try:
        expected = {(name,row):sheet.cell(row,1).value for name in ('필드확인','검수확인')
            for sheet in (book[name],) for row in range(5,sheet.max_row+1) if sheet.cell(row,1).value}
        seen = set()
        for item in result['items']:
            key = (item['sheet'],item['row'])
            if key in seen or expected.get(key)!=item['id'] or item.get('reviewer')!=reviewer:
                raise EvaluationNotReady('사람 확인의 시트·행·ID·검수자가 다릅니다.')
            if (item.get('decision')!='확인' and item['id'] not in wording_resolutions) or not str(item.get('reason','')).strip():
                raise EvaluationNotReady('수정·보류 항목은 별도로 원문을 재확인한 뒤 반영해야 합니다.')
            moment = datetime.fromisoformat(item['reviewed_at'].replace('Z','+00:00'))
            if moment.tzinfo is None:
                raise EvaluationNotReady('사람 확인 시각에 시간대가 필요합니다.')
            seen.add(key)
            book[key[0]].cell(key[1],11,item['decision'])
            book[key[0]].cell(key[1],12,item['reason'])
        if seen!=set(expected):
            raise EvaluationNotReady('사람 확인 항목이 누락됐습니다.')
        book['안내']['B12']=reviewer
        p['human'].mkdir(parents=True)
        reviewed=p['human']/'human_reviewed.xlsx'
        book.save(reviewed)
    finally:
        book.close()
    original=p['human']/result_path.name
    shutil.copyfile(result_path,original)
    resolution_copy=None
    if resolution_path is not None:
        resolution_copy=p['human']/'question_wording_resolution.json'
        shutil.copyfile(resolution_path,resolution_copy)
    labels=import_review(reviewed,p['dataset'],p['comparison']/'comparison.json',p['blind']/'replies',reviewer,original,
        split='test',resolution_path=resolution_copy)
    return {'human_confirmed':True,'human_direct_count':labels['review']['human_direct_count'],
            'ai_agreement_count':labels['review']['ai_agreement_count'],
            'prompt_wording_error_count':labels['review'].get('prompt_wording_error_count',0),
            'original_human_decision_counts':labels['review'].get('original_human_decision_counts'),
            'final_inference_started':False}


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    group=parser.add_mutually_exclusive_group()
    group.add_argument('--replies',type=Path)
    group.add_argument('--human-result',type=Path)
    parser.add_argument('--wording-resolution',type=Path,help='원본 답을 바꾸지 않는 질문 문구 오류의 근거 기록')
    args=parser.parse_args(argv)
    try:
        if not args.replies and not args.human_result:
            update_blind_prompt(PROJECT)
        if args.wording_resolution and not args.human_result:
            raise EvaluationNotReady('질문 정정에는 --human-result가 필요합니다.')
        report=reconcile(PROJECT,args.replies) if args.replies else confirm(PROJECT,args.human_result,args.wording_resolution) if args.human_result else readiness(PROJECT)
        write_json(paths(PROJECT)['output']/'readiness.json',report)
        print(json.dumps(report,ensure_ascii=False,indent=2))
        return 0
    except (EvaluationNotReady,ValueError,OSError,KeyError,TypeError) as exc:
        print('test2 준비 중단: '+str(exc),file=sys.stderr)
        return 2


if __name__=='__main__':
    raise SystemExit(main())
