# IDs encode the contract; no app predictions determine the questions.
# Text from independent review must not become an Excel formula.
"""Prepare heldout labels and blind review, without predictions or final attempts."""
import argparse
import copy
import importlib.util
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))
from backend.evaluation_core import EvaluationNotReady, file_hash, verify_dataset, verify_freeze
from backend.evaluation_release import implementation_hashes, verify_release, check_final_attempt
from backend.scripts.build_reference_labels import build
from backend.scripts.package_independent_review import package
from backend.scripts.compare_reference_labels import main as compare_main
from backend.scripts.import_reconciled_labels import import_review
from backend.evaluation_review import compare_labels, merge_independent_parts, validate_reference_draft

def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

def freeze_review_protocol(project=PROJECT):
    parent = project / 'docs/eval/dev/guard_v3/release_lock.json'
    path = project / 'docs/eval/test/preparation/release_lock.json'
    base = json.loads(parent.read_text(encoding='utf-8'))
    current = implementation_hashes(project)
    changes = sorted((k for k in set(current) | set(base['implementation_sha256']) if current.get(k) != base['implementation_sha256'].get(k)))
    if changes != ['backend/evaluation_core.py']:
        raise EvaluationNotReady('정답 검수 절차 외 고정 구현 변경이 있습니다. AI를 변경한 상태로 test를 준비하지 않습니다.')
    record = copy.deepcopy(base)
    record.update(implementation_sha256=current, frozen_at=datetime.now(timezone.utc).isoformat(), parent_release_sha256=file_hash(parent), parent_release='docs/eval/dev/guard_v3/release_lock.json', review_protocol_change='test에도 dev와 같은 표본 확인 출처 검증 허용. 추출·판정·모델·프롬프트·가드는 변경 없음', review_protocol_changed_files=changes)
    verify_release(project, record, record['selected_model'])
    if path.exists():
        existing = json.loads(path.read_text(encoding='utf-8'))
        if existing['implementation_sha256'] != current or existing['parent_release_sha256'] != file_hash(parent):
            raise EvaluationNotReady('이미 고정한 test 준비 버전을 덮어쓰지 않습니다.')
        verify_release(project, existing, existing['selected_model'])
    else:
        save_json(path, record)
    return path

def page_builder():
    spec = importlib.util.spec_from_file_location('human_review_builder', PROJECT / 'eval/human_review_page/build_review_page.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def review_questions(reference):
    questions = {}
    first_cid = min(reference['contracts'], key=int)

    def add(rid, reason):
        if rid in questions:
            questions[rid]['reason'] += ' ' + reason
        else:
            questions[rid] = {'review_id': rid, 'reason': reason}

    for cid, contract in reference['contracts'].items():
        for did, doc in reference['documents'].items():
            if int(did) // 100 != int(cid):
                continue
            kind = doc['document_type']
            if kind in {'inspection_confirmation', 'contract', 'estimate'}:
                reason = {'inspection_confirmation': '검수확인서는 청구서가 아니어도 계약 전체 금액이면 FULL입니다. dev에서 확정한 같은 기준을 원문으로 확인합니다.',
                          'contract': '계약서 자체의 전체 계약 금액이면 FULL입니다. 일부 지급액과 구분해 원문으로 확인합니다.',
                          'estimate': '견적서의 전체 견적 금액이면 FULL입니다. 적용 범위가 표시되지 않은 참고값은 REVIEW입니다.'}[kind]
                add(f'{cid}/{did}/amount_scope', reason)
            scope = doc['fields']['amount_scope']
            if kind == 'tax_invoice' and scope.get('value') == 'PARTIAL':
                add(f'{cid}/{did}/amount_scope', "'분할 청구 선금 또는 잔금'은 어느 회차인지 몰라도 일부 청구(PARTIAL)입니다.")
            if kind == 'estimate' and doc['fields']['tax_amount']['state'] == 'REVIEW':
                add(f'{cid}/{did}/tax_amount', '참고값 A/B의 뜻을 확정할 수 없으면 세액도 REVIEW입니다. 세액이 명확히 없는 ABSENT와 구분합니다.')
        for check in contract['checks']:
            rid = f"{cid}/contract/{check['rule_id']}/{check['field']}"
            if check['rule_id'] == 'BUSINESS_STATUS':
                add(rid, '국세청 미연동 평가는 영업 상태를 추정하지 않고 항상 REVIEW입니다. dev에서 확정한 같은 기준입니다.')
            if check['status'] == 'FAIL' and check['rule_id'] in {'AMOUNT_COMPONENTS', 'AMOUNT_MATCH_TAX_INVOICE'}:
                add(rid, '원문 표시 합계, 공급가액+세액, 전액 청구와 등록 금액을 대조합니다. 전액 청구가 명시되고 합계가 다르면 FAIL입니다.')
            # Use real labelled estimate/invoice rows for the applicability topic;
            # contract/inspection candidate checks are not scoring-label slots.
            doc = reference['documents'].get(check['field'], {})
            expected_kind = {'AMOUNT_MATCH_ESTIMATE': 'estimate', 'AMOUNT_MATCH_TAX_INVOICE': 'tax_invoice'}.get(check['rule_id'])
            if cid == first_cid and expected_kind and doc.get('document_type') == expected_kind:
                add(rid, '금액 대조는 견적서·세금계산서에 적용합니다. 계약서·검수확인서의 FULL은 금액 범위 뜻이며 같은 대조 검사를 적용한다는 뜻이 아닙니다. 원문 금액과 등록 합계를 비교하세요.')
    return list(questions.values())

def prepare(root, font):
    check_final_attempt(root)
    release = freeze_review_protocol()
    reference_path = root / 'reference_labels_test.json'
    reference = json.loads(reference_path.read_text(encoding='utf-8')) if reference_path.exists() else build(root, font, split='test')
    manifest = json.loads((root / 'manifest.json').read_text(encoding='utf-8'))
    verify_dataset(root, manifest, split='test')
    validate_reference_draft(manifest, reference, 'test')
    verify_freeze(root / 'manifest.json', reference_path)
    destination = PROJECT / 'eval/claude_blind_review_test'
    if not destination.exists():
        package(root, destination, split='test')
    if not destination.with_suffix('.zip').is_file() or not (destination / 'package_receipt.json').is_file():
        raise EvaluationNotReady('독립 판독 묶음이 불완전합니다. 기존 폴더를 덮어쓰지 않습니다.')
    receipt = json.loads((destination / 'package_receipt.json').read_text(encoding='utf-8'))
    for name, digest in receipt['files'].items():
        path = (destination / name).resolve()
        if not path.is_relative_to(destination.resolve()) or file_hash(path) != digest:
            raise EvaluationNotReady('독립 판독 묶음의 수령 해시가 달라졌습니다.')
    (destination / 'replies').mkdir(exist_ok=True)
    questions = PROJECT / 'eval/human_questions_test_v1.json'
    if not questions.exists():
        save_json(questions, review_questions(reference))
    page_builder().build_page(split='test', dataset=root, blind=destination)
    return release

def format_review_guide(guide):
    from openpyxl.styles import Font, Alignment
    for cell, label in {'A2': '자료 버전', 'A3': '대조 결과 해시', 'A4': '원문 목록 해시',
                        'A5': '사람 확인할 칸', 'A6': '일치 표본 비율', 'A13': '평가 구분'}.items():
        guide[cell] = label
    guide.column_dimensions['A'].width = 26
    guide.column_dimensions['B'].width = 94
    guide['A1'].font = Font(bold=True, size=13)
    guide.row_dimensions[1].height = 28
    guide['B6'].number_format = '0%'
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top')


def workbook_from_comparison(root, comparison_path, output):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    if output.exists():
        raise EvaluationNotReady('검수 양식을 덮어쓰지 않습니다.')
    comparison = json.loads(comparison_path.read_text(encoding='utf-8'))
    if comparison['split'] != 'test' or not comparison['review']['independent_complete']:
        raise EvaluationNotReady('test 독립 판독이 모두 완료되어야 사람 확인 목록을 만들 수 있습니다.')
    book = openpyxl.Workbook()
    guide = book.active
    guide.title = '안내'
    guide['A1'] = 'test 원문 확인 — 사람 직접 확인과 AI 일치를 구분'
    for cell, value in {'B2': comparison['dataset_version'], 'B3': file_hash(comparison_path), 'B4': file_hash(root / 'manifest.json'), 'B5': comparison['required_human_count'], 'B6': comparison['sampling']['rate'], 'B13': 'test'}.items():
        guide[cell] = value
    guide['A12'] = '검수자'
    format_review_guide(guide)
    headers = ['ID', '원문', '확인할 항목', 'Codex 상태', 'Codex 값', 'Claude 상태', 'Claude 값', '선정 이유', '최종 상태', '최종 값', '사람 확인', '근거', '해석 기준']
    sheets = {name: book.create_sheet(name) for name in ('필드확인', '검수확인')}
    for sheet in sheets.values():
        sheet.append(['원문과 I·J열의 제안 답을 비교하고, K·L열은 직접 확인한 뒤 입력하세요.'])
        sheet.append([])
        sheet.append([])
        sheet.append(headers)
        sheet.freeze_panes = 'D5'
        sheet.auto_filter.ref = 'A4:M4'
        for cell in sheet[4]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='334155')
        for column, width in {'A': 42, 'B': 48, 'C': 30, 'D': 18, 'E': 24, 'F': 18, 'G': 24, 'H': 20, 'I': 18, 'J': 24, 'K': 16, 'L': 44, 'M': 48}.items():
            sheet.column_dimensions[column].width = width
    reasons = {'RANDOM_SAMPLE': '일치 표본', 'INTERPRETATION': '해석 확인', 'AMBIGUITY': '모호한 값', 'DISAGREEMENT': '불일치'}

    def state_value(row, value):
        if isinstance(value, dict):
            answer = value.get('value')
            if row['key'] == 'business_registration_no' and isinstance(answer, str) and len(answer) == 10 and answer.isdigit():
                answer = f'{answer[:3]}-{answer[3:5]}-{answer[5:]}'
            return (value['state'], answer)
        if row['key'] == 'document_type':
            return ('TYPE', value)
        if row['key'] == 'missing_types':
            return ('MISSING_TYPES', json.dumps(value, ensure_ascii=False))
        return (value, None)
    for row in comparison['rows']:
        if not row['human_review_required']:
            continue
        a, b = (state_value(row, row['reference']), state_value(row, row['independent']))
        sheet = sheets['필드확인' if row['document_id'] is not None else '검수확인']
        sheet.append([row['review_id'], str(root / row['source_path']) if row['source_path'] else '등록 계약과 전체 문서', row['key'], *a, *b, reasons[row['selection_reason']], *a, None, None, row.get('human_question', '')])
        if row['key'] == 'business_registration_no':
            sheet.cell(sheet.max_row, 10).number_format = '@'
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = 's'
        for cell in sheet[sheet.max_row][8:12]:
            cell.fill = PatternFill('solid', fgColor='FEF9C3')
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    return output

def reconcile(root, replies):
    check_final_attempt(root)
    freeze_review_protocol()
    manifest = json.loads((root / 'manifest.json').read_text(encoding='utf-8'))
    reference = json.loads((root / 'reference_labels_test.json').read_text(encoding='utf-8'))
    files = sorted(replies.glob('*.json'))
    independent = merge_independent_parts([json.loads(p.read_text(encoding='utf-8-sig')) for p in files], split='test')
    if not compare_labels(manifest, reference, independent)['review']['independent_complete']:
        raise EvaluationNotReady('Claude의 모든 원문 판독과 답을 먼저 완료하세요. 불완전한 확인 목록은 만들지 않습니다.')
    # Preserve the draft and blind replies. Version only the review questions;
    # the new interpretation topics are human confirmation inputs, not scoring changes.
    questions_path = PROJECT / 'eval/human_questions_test_v2.json'
    questions = review_questions(reference)
    if questions_path.exists():
        if json.loads(questions_path.read_text(encoding='utf-8')) != questions:
            raise EvaluationNotReady('기존 test 질문 버전을 덮어쓰지 않습니다.')
    else:
        save_json(questions_path, questions)
    destination = PROJECT / 'eval/reconciliation-test-v1'
    if destination.exists() or (PROJECT / 'eval/human_review_test.xlsx').exists():
        raise EvaluationNotReady('이미 만든 test 검수 목록을 덮어쓰지 않습니다.')
    if compare_main([*[str(p) for p in files], '--split', 'test', '--dataset', str(root), '--output', str(destination), '--sample-rate', '0.10', '--sample-method', 'case_document_type', '--human-questions', str(questions_path)]) != 0:
        raise EvaluationNotReady('독립 판독 대조를 완료하지 못했습니다.')
    workbook = workbook_from_comparison(root, destination / 'comparison.json', PROJECT / 'eval/human_review_test.xlsx')
    page_builder().build_page(split='test', dataset=root, workbook=workbook)

def confirm(root, result_path):
    import openpyxl
    source = PROJECT / 'eval/human_review_test.xlsx'
    result = json.loads(result_path.read_text(encoding='utf-8-sig'))
    if result.get('split') != 'test' or result.get('source_xlsx_sha256') != file_hash(source):
        raise EvaluationNotReady('검수 결과의 test 구분 또는 원본 양식 해시가 다릅니다.')
    reviewer = str(result.get('reviewer', '')).strip()
    if not reviewer or result_path.suffix.lower() != '.json':
        raise EvaluationNotReady('검수자와 원본 JSON 파일이 필요합니다.')
    directory = PROJECT / 'eval/human_reviews/test'
    reviewed = directory / 'human_reviewed.xlsx'
    if directory.exists():
        raise EvaluationNotReady('이전 사람 검수 출처를 덮어쓰지 않습니다.')
    book = openpyxl.load_workbook(source)
    try:
        expected = {(name, n): sheet.cell(n, 1).value for name in ('필드확인', '검수확인') for sheet in [book[name]] for n in range(5, sheet.max_row + 1) if sheet.cell(n, 1).value}
        seen = set()
        for item in result['items']:
            key = (item['sheet'], item['row'])
            if key in seen or expected.get(key) != item['id'] or item['reviewer'] != reviewer:
                raise EvaluationNotReady('사람 확인의 시트·행·ID·검수자가 다릅니다.')
            seen.add(key)
            if item.get('decision') != '확인' or not str(item.get('reason', '')).strip():
                raise EvaluationNotReady('수정·보류 항목은 원문과 정답을 먼저 재확인하세요. 확정하지 않았습니다.')
            timestamp = datetime.fromisoformat(item['reviewed_at'].replace('Z', '+00:00'))
            if timestamp.tzinfo is None:
                raise EvaluationNotReady('사람 확인 시각에 시간대가 필요합니다.')
            book[item['sheet']].cell(item['row'], 11, item['decision'])
            book[item['sheet']].cell(item['row'], 12, item['reason'])
        if seen != set(expected):
            raise EvaluationNotReady('사람 확인 목록이 누락되었습니다.')
        book['안내']['B12'] = reviewer
        directory.mkdir(parents=True)
        book.save(reviewed)
    finally:
        book.close()
    original = directory / result_path.name
    shutil.copyfile(result_path, original)
    return import_review(reviewed, root, PROJECT / 'eval/reconciliation-test-v1/comparison.json', PROJECT / 'eval/claude_blind_review_test/replies', reviewer, original, split='test')

def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dataset', type=Path, default=PROJECT / 'eval/dataset')
    parser.add_argument('--font', type=Path, default=Path('C:/Windows/Fonts/malgun.ttf'))
    group = parser.add_mutually_exclusive_group()
    group.add_argument('--replies', type=Path)
    group.add_argument('--human-result', type=Path)
    args = parser.parse_args(argv)
    try:
        root = args.dataset.resolve()
        if args.replies:
            reconcile(root, args.replies)
        elif args.human_result:
            confirm(root, args.human_result)
        else:
            prepare(root, args.font)
        print('test 정답 준비 처리 완료. 이 명령은 성능 측정이나 최종 평가를 실행하지 않습니다.')
        return 0
    except (EvaluationNotReady, OSError, ValueError, KeyError, TypeError) as exc:
        print('test 준비 중단: ' + str(exc), file=sys.stderr)
        return 2
if __name__ == '__main__':
    raise SystemExit(main())
